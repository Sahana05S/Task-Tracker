"""Reconcile the app's September 2026 ten-factor scores with the workbook's values.

Reads the tracker workbook (people through tools/seed_sheet_to_project.py's read_sheet;
per-employee values from '1_2 Performance Analysis' rows 3+ and the raw/median cells of
'B_Summary & Analysis'), then reads the app's values for the same people through the code
the live report uses: public.performance_source_context as the director, evaluated by the
scoring module's parsePerformanceSources + resolvePerformance (tools/reconcile_app_scores.mjs).
Prints one table per employee x factor (sheet, app, difference, app unavailable reason),
compares at display precision (factors 6 decimals, scores 2, half away from zero) and
classifies every difference. With --report it writes the Markdown report.

Sources:
  --target local   a disposable loopback PostgreSQL (default). --build creates it from a
                   migrations folder and seeds the sheet with tools/seed_sheet_to_project.py.
  --target hosted  a project through the Supabase CLI query channel, READ ONLY transactions
                   only; use it only with the user's go-ahead for that project.
Scoring rules are never changed here; differences are classified, never hidden.

  python tools/reconcile_sheet_scores.py --self-test
"""

import argparse
import datetime as dt
import importlib.util
import json
import re
import subprocess
import sys
import tempfile
import uuid
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FACTORS = ("productivity", "criticality", "loading", "completion", "progress",
           "sales", "facts", "ideas", "suggestions", "freshness")
COEFFICIENTS = {"productivity": "1.5", "criticality": "0.5", "loading": "40", "completion": "1.5", "progress": "0.0125",
                "sales": "1", "facts": "1", "ideas": "1", "suggestions": "1", "freshness": "1"}
FACTOR_PLACES, SCORE_PLACES = 6, 2
PERF_SHEET, SUMMARY_SHEET = "1_2 Performance Analysis", "B_Summary & Analysis"
PERF_COLUMNS = dict(zip(FACTORS, "BCDEFGHIJK"))  # 1_2 Performance Analysis: differences with median; L = Sum
# B_Summary & Analysis median cells; freshness has none (the workbook adds its raw sum).
MEDIAN_CELLS = {"productivity": "L2", "progress": "M2", "criticality": "AA2", "loading": "AK2", "completion": "AR2",
                "facts": "AW2", "ideas": "BD2", "suggestions": "BJ2", "sales": "BP2"}
# Summary blocks whose rows carry an unlabeled raw value next to its median difference; the
# employee-name columns are dynamic arrays the xlsx export does not keep, so a row is tied to
# a person by its difference matching exactly one labelled 1_2 row.
SUMMARY_BLOCKS = {"productivity": ("K", "L", 4), "progress": (None, "M", 4),
                  "criticality": ("AA", "AB", 5), "loading": ("AI", "AJ", 4)}
SUMMARY_LAST_ROW = 51
ATTENDANCE_COLUMN = "J"  # hours from the external attendance workbook, same rows as productivity

# Difference classes. The text is what the report says about each one.
CLASSES = {
    "median_cohort": "Median cohort: the app's median is over all active members (measured zero included, "
                     "D04); the sheet's median is over the members listed in that block, and 1_2 shows 0 "
                     "(IFERROR) for anyone missing from it.",
    "as_of_window": "As-of window: the app reads observations as of the chosen as-of instant (task states, "
                    "freshness and carried workload on that date); the sheet was computed at its own date.",
    "logs_outside_period": "Logs dated outside September: the app's progress is September-earned credit (D02), "
                           "and the seed's last-log mode puts each task's whole sheet progress on its latest-dated "
                           "log; where that log is October-dated (88 such logs; 5 more are dated 2006) the credit "
                           "falls outside September. The sheet counts the tasks' whole scaled progress. The "
                           "residual check: sheet raw = app raw + credit on outside-dated logs.",
    "attendance": "Attendance validity (M06): no September attendance is entered in the app, so attendance "
                  "hours (the productivity denominator) and ranking eligibility are unavailable; the sheet "
                  "imports hours from the external attendance workbook.",
    "sales": "Sales flag counts (D05): the app counts BN/BO/BP flags on eligible sales_logs entries; the "
             "sheet's sales block is empty.",
    "sheet_source_empty": "Sheet source empty: the workbook's input block for this factor has no rows "
                          "(median #NUM! or an empty block), so 1_2 shows 0 through IFERROR for everyone while the "
                          "app measures a value and centres it on its median. For freshness the app's value also "
                          "depends on the as-of date (task states on that day).",
    "import_rounding": "Import rounding: the seed stores task progress at two decimals (as the app does), the "
                       "sheet keeps it unrounded; the residual equals the member's scale x rounding gap exactly.",
    "UNEXPLAINED": "UNEXPLAINED: none of the known causes accounts for the difference.",
}
KNOWN_OUTSIDE_LOGS = {"october": 88, "2006": 5}
# Sheet medians are cached to about 9 significant digits (M2 2001.193269), so raw = median +
# difference carries up to 5e-7 of that rounding; residuals below this are equal.
TOLERANCE = Decimal("0.000001")
FRESHNESS_BLOCK = ("CA", 5)


# --------------------------------------------------------------------------- pure helpers

def dec(value):
    if value is None or isinstance(value, Decimal):
        return value
    if isinstance(value, float):
        return Decimal(repr(value))
    if isinstance(value, (int, str)):
        try:
            return Decimal(str(value))
        except ArithmeticError:
            return None
    if isinstance(value, dict) and "numerator" in value and "denominator" in value:  # Exact json
        return Decimal(str(value["numerator"])) / Decimal(str(value["denominator"]))
    return None


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def rounded(value, places):
    """Display rounding: half away from zero (Decimal ROUND_HALF_UP), as the app displays."""
    value = dec(value)
    return None if value is None else value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)


def fmt(value, places):
    value = rounded(value, places)
    if value is None:
        return "-"
    return format(value + 0, "f")  # +0 turns -0.000000 into 0.000000


def compare(sheet, app, places):
    """('match'|'differs'|'app_unavailable'|'sheet_missing', difference app-sheet at display precision)."""
    s, a = rounded(sheet, places), rounded(app, places)
    if a is None:
        return "app_unavailable", None
    if s is None:
        return "sheet_missing", None
    diff = a - s
    return ("match" if diff == 0 else "differs"), diff


def sheet_values(perf, summary, people):
    """Per-employee sheet values from cell dicts {coordinate: value}.

    perf: '1_2 Performance Analysis'; summary: 'B_Summary & Analysis'; people: read_sheet people.
    Returns {employee_id: {name, score, attendance, factors: {f: {delta, raw, median, in_cohort, source}}}}.
    """
    by_name = {p["name"]: p for p in people}
    rows = {}
    for r in range(3, 3 + 21):
        name = perf.get(f"A{r}")
        if not isinstance(name, str) or not name.strip():
            break
        person = by_name.get(name.strip())
        if person is None:
            raise ValueError(f"{PERF_SHEET}!A{r} {name!r} is not in C_Team Members")
        rows[r] = person
    if not rows:
        raise ValueError(f"{PERF_SHEET} has no employee rows")
    medians = {f: summary.get(cell) if is_number(summary.get(cell)) else None for f, cell in MEDIAN_CELLS.items()}
    errors = {f: summary.get(cell) for f, cell in MEDIAN_CELLS.items() if isinstance(summary.get(cell), str)}
    out = {}
    for r, person in rows.items():
        factors = {}
        for f in FACTORS:
            delta = perf.get(f"{PERF_COLUMNS[f]}{r}")
            factors[f] = {"delta": delta if is_number(delta) else None, "raw": None, "median": medians.get(f),
                          "in_cohort": None, "source": errors.get(f) and f"median {MEDIAN_CELLS[f]} {errors[f]}"}
        score = perf.get(f"L{r}")
        out[person["employee_id"]] = {"name": person["name"], "row": r, "score": score if is_number(score) else None,
                                      "attendance": None, "factors": factors}
    col, first_row = FRESHNESS_BLOCK
    if not any(is_number(summary.get(f"{col}{r}")) for r in range(first_row, SUMMARY_LAST_ROW + 1)):
        for v in out.values():
            v["factors"]["freshness"]["source"] = f"freshness block {col}{first_row}:{col}{SUMMARY_LAST_ROW} empty"
    for f, (raw_col, delta_col, first) in SUMMARY_BLOCKS.items():
        listed = 0
        for r in range(first, SUMMARY_LAST_ROW + 1):
            delta = summary.get(f"{delta_col}{r}")
            if not is_number(delta):
                continue
            listed += 1
            matches = [e for e, v in out.items() if v["factors"][f]["delta"] is not None
                       and abs(dec(v["factors"][f]["delta"]) - dec(delta)) < Decimal("1e-9")]
            if len(matches) != 1:
                continue  # ambiguous or unmatched rows tie to nobody; the person stays "not proven"
            entry = out[matches[0]]["factors"][f]
            raw = summary.get(f"{raw_col}{r}") if raw_col else None
            if raw_col is None and medians.get(f) is not None:
                raw = float(dec(medians[f]) + dec(delta))
            entry.update(raw=raw if is_number(raw) else None, in_cohort=True, source=f"{SUMMARY_SHEET}!{delta_col}{r}")
            if f == "productivity":
                hours = summary.get(f"{ATTENDANCE_COLUMN}{r}")
                out[matches[0]]["attendance"] = hours if is_number(hours) else None
        for v in out.values():
            entry = v["factors"][f]
            if entry["in_cohort"] is None and listed:
                entry["in_cohort"] = False
                entry["source"] = f"not listed in {SUMMARY_SHEET} ({listed} rows); 1_2 IFERROR shows {entry['delta']}"
            entry["cohort_size"] = listed
    return out


def classify(factor, sheet, app, evidence):
    """Classes for one employee x factor; each class is backed by a check on the numbers.

    sheet/app are the factor entries; evidence is per member: outside_credit (scaled progress
    on logs dated outside the period), sheet_task_scaled / imported_task_scaled (the member's
    tasks' scale x progress as the sheet holds it and as the seed stored it, 2 decimals).
    """
    status, _ = compare(sheet["delta"], app.get("delta"), FACTOR_PLACES)
    if status == "match":
        return ["match"]
    reasons = app.get("reasons") or []
    classes = []
    if status == "app_unavailable":
        if any("attendance" in r for r in reasons):
            classes.append("attendance")
        if "task_history_before_observed_baseline" in reasons or evidence.get("history") is False:
            classes.append("as_of_window")
        if factor == "sales" and any(r.startswith("sale") or "legacy" in r or "unattributed" in r for r in reasons):
            classes.append("sales")
        if factor == "productivity" and classes and dec(evidence.get("outside_credit") or 0) > 0:
            classes.append("logs_outside_period")  # the numerator would differ too once attendance exists
        return classes or ["UNEXPLAINED"]
    raw_status, _ = compare(sheet.get("raw"), app.get("raw"), FACTOR_PLACES)
    sheet_empty = bool(sheet.get("source") and ("#NUM!" in sheet["source"] or "empty" in sheet["source"]))
    if sheet_empty:
        classes.append("sheet_source_empty")
    elif sheet.get("in_cohort") is False:
        # Not in the sheet's block: 1_2 shows IFERROR 0, the app measures a value and centres it.
        classes.append("median_cohort")
    elif raw_status == "match":
        # Same raw value, different difference: only the median can differ; it must be a cohort difference.
        if sheet.get("cohort_size") != app.get("population"):
            classes.append("median_cohort")
    elif factor in ("progress", "productivity") and sheet.get("raw") is not None and app.get("raw") is not None:
        outside = dec(evidence.get("outside_credit") or 0)
        residual = dec(sheet["raw"]) - (dec(app["raw"]) + outside)
        gap = dec(evidence.get("sheet_task_scaled") or 0) - dec(evidence.get("imported_task_scaled") or 0)
        if outside > 0:
            classes.append("logs_outside_period")
        if factor == "progress" and abs(residual) >= TOLERANCE:
            classes.append("import_rounding" if gap and abs(residual - gap) < TOLERANCE else "UNEXPLAINED")
        if sheet.get("cohort_size") != app.get("population"):
            classes.append("median_cohort")
    elif factor == "sales" and evidence.get("sales_entries"):
        classes.append("sales")
    return list(dict.fromkeys(classes)) or ["UNEXPLAINED"]


def median_shift(cells):
    """Second pass over one factor's cells (all employees). A cell whose raw value matches, in a
    cohort of the same size, can only differ through the median; it inherits the explained
    classes of the cells whose raw values differ (marked 'via median'). Returns new class lists."""
    sources = [c for c in cells if compare(c["sheet"].get("raw"), c["app"].get("raw"), FACTOR_PLACES)[0] == "differs"
               and "UNEXPLAINED" not in c["classes"]]
    inherited = list(dict.fromkeys(k for c in sources for k in c["classes"] if k != "median_cohort"))
    out = []
    for c in cells:
        same_raw = compare(c["sheet"].get("raw"), c["app"].get("raw"), FACTOR_PLACES)[0] == "match"
        same_cohort = c["sheet"].get("cohort_size") == c["app"].get("population")
        if c["classes"] == ["UNEXPLAINED"] and same_raw and same_cohort and inherited:
            out.append([f"{k} (via median)" for k in inherited])
        else:
            out.append(c["classes"])
    return out


def sheet_score_check(entry):
    """The workbook's own L column from its 1_2 values (first five weighted, last five added)."""
    total = Decimal(0)
    for f in FACTORS:
        value = entry["factors"][f]["delta"]
        total += dec(value or 0) * Decimal(COEFFICIENTS[f])
    return total


def self_test():
    people = [{"name": "A", "employee_id": "E1"}, {"name": "B", "employee_id": "E2"}, {"name": "C", "employee_id": "E3"}]
    perf = {"A3": "A", "B3": -2.0, "C3": 0, "D3": 0, "E3": 0, "F3": -10.0, "G3": 0, "H3": 0, "I3": 0, "J3": 0, "K3": 0,
            "A4": "B", "B4": 0.0, "C4": 1.0, "D4": 0.1, "E4": 0, "F4": 0.0, "G4": 0, "H4": 0, "I4": 0, "J4": 0, "K4": 0,
            "A5": "C", "B5": 5.0, "C5": -1.0, "D5": -0.1, "E5": 0, "F5": 30.0, "G5": 0, "H5": 0, "I5": 0, "J5": 0, "K5": 0}
    for r, e in ((3, "A"), (4, "B"), (5, "C")):
        perf[f"L{r}"] = float(sheet_score_check({"factors": {f: {"delta": perf[f"{PERF_COLUMNS[f]}{r}"]} for f in FACTORS}}))
    summary = {"L2": 2.0, "M2": 10.0, "AA2": 2.0, "AK2": 0.5, "AR2": "#NUM!",
               "J4": 0, "K4": 0, "L4": -2.0, "M4": -10.0, "J5": 10, "K5": 7.0, "L5": 5.0, "M5": 30.0,
               "J6": 8, "K6": 2.0, "L6": 0.0, "M6": 0.0,
               "AA5": 3.0, "AB5": 1.0, "AA6": 1.0, "AB6": -1.0, "AI4": 0.6, "AJ4": 0.1, "AI5": 0.4, "AJ5": -0.1}
    got = sheet_values(perf, summary, people)
    assert got["E3"]["factors"]["productivity"]["raw"] == 7.0 and got["E3"]["attendance"] == 10
    assert got["E1"]["factors"]["progress"]["raw"] == 0.0 and got["E3"]["factors"]["progress"]["raw"] == 40.0
    assert got["E2"]["factors"]["criticality"]["raw"] == 3.0 and got["E1"]["factors"]["criticality"]["in_cohort"] is False
    assert got["E1"]["factors"]["criticality"]["cohort_size"] == 2
    assert "#NUM!" in got["E1"]["factors"]["completion"]["source"]
    assert dec(got["E3"]["score"]) == sheet_score_check(got["E3"])
    # Ambiguous rows (two people with the same difference) tie to nobody.
    tie = sheet_values(dict(perf, B4=5.0), dict(summary, L6=5.0), people)
    assert tie["E2"]["factors"]["productivity"]["raw"] is None and tie["E3"]["factors"]["productivity"]["raw"] is None
    try:
        sheet_values(dict(perf, A4="Nobody"), summary, people)
        raise AssertionError("unknown name accepted")
    except ValueError:
        pass
    # Display precision: half away from zero, both signs; 6 places for factors, 2 for scores.
    assert fmt("0.0000005", 6) == "0.000001" and fmt("-0.0000005", 6) == "-0.000001" and fmt("-0.0000004", 6) == "0.000000"
    assert fmt("2.675", 2) == "2.68" and fmt(None, 2) == "-" and fmt({"numerator": "1", "denominator": "3"}, 6) == "0.333333"
    assert compare(1.0000004, 1.0000001, 6) == ("match", Decimal("0E-6"))
    assert compare(1.0, None, 6)[0] == "app_unavailable" and compare(None, 1.0, 6)[0] == "sheet_missing"
    assert compare(33.084, 33.09, 2) == ("differs", Decimal("0.01"))
    scaled = task_scaled([{"tag": "X", "scale": 100, "progress": 4.1038, "assignee": "A"},
                          {"tag": "Y", "scale": 50, "progress": 10, "assignee": "A"},
                          {"tag": "Z", "scale": 50, "progress": 10, "assignee": "Nobody"}], people,
                         lambda v: dec(v).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    assert scaled == {"E1": [Decimal("9.103800"), Decimal("9.1000")]}
    assert "empty" in got["E1"]["factors"]["freshness"]["source"]
    assert "empty" not in (sheet_values(perf, dict(summary, CA5=3), people)["E1"]["factors"]["freshness"]["source"] or "")
    # Classification: each class needs its numeric evidence.
    sheet = {"delta": 1.0, "raw": 3.0, "median": 2.0, "in_cohort": True, "cohort_size": 2, "source": "x"}
    assert classify("criticality", sheet, {"delta": 1.0, "raw": 3.0, "median": 2.0, "population": 2}, {}) == ["match"]
    assert classify("criticality", sheet, {"delta": 2.0, "raw": 3.0, "median": 1.0, "population": 4}, {}) == ["median_cohort"]
    # Different raw values with no evidence are never put down to the cohort.
    assert classify("criticality", sheet, {"delta": 2.5, "raw": 3.5, "median": 1.0, "population": 4}, {}) == ["UNEXPLAINED"]
    assert classify("criticality", sheet, {"delta": 2.0, "raw": 3.0, "median": 1.0, "population": 2}, {}) == ["UNEXPLAINED"]
    absent = dict(sheet, delta=0, raw=None, in_cohort=False)
    assert classify("loading", absent, {"delta": -0.5, "raw": 0, "median": 0.5, "population": 4}, {}) == ["median_cohort"]
    assert classify("productivity", sheet, {"delta": None, "reasons": ["attendance_denominator_missing"]}, {}) == ["attendance"]
    assert classify("productivity", sheet, {"delta": None, "reasons": ["attendance_denominator_missing"]},
                    {"outside_credit": 5}) == ["attendance", "logs_outside_period"]
    progress = {"delta": 1.0, "raw": 6.0, "median": 5.0, "in_cohort": True, "cohort_size": 4, "source": "x"}
    assert classify("progress", progress, {"delta": 0.0, "raw": 1.0, "median": 1.0, "population": 4},
                    {"outside_credit": "5"}) == ["logs_outside_period"]
    assert classify("progress", progress, {"delta": 0.0, "raw": 1.0, "median": 1.0, "population": 4},
                    {"outside_credit": "4.9", "sheet_task_scaled": "10.1", "imported_task_scaled": "10"}) ==         ["logs_outside_period", "import_rounding"]
    assert classify("progress", progress, {"delta": 0.0, "raw": 1.0, "median": 1.0, "population": 4},
                    {"outside_credit": "4", "sheet_task_scaled": "10.1", "imported_task_scaled": "10"}) ==         ["logs_outside_period", "UNEXPLAINED"]
    assert classify("progress", progress, {"delta": 0.0, "raw": 1.0, "median": 1.0, "population": 4}, {}) == ["UNEXPLAINED"]
    assert classify("criticality", sheet, {"delta": None, "reasons": ["task_history_before_observed_baseline"]}, {}) == ["as_of_window"]
    assert classify("sales", sheet, {"delta": None, "reasons": ["sales_legacy_unreconciled"]}, {}) == ["sales"]
    assert classify("ideas", sheet, {"delta": None, "reasons": ["mystery"]}, {}) == ["UNEXPLAINED"]
    empty = {"delta": 0, "raw": None, "median": None, "in_cohort": None, "source": "median AR2 #NUM!"}
    assert classify("completion", empty, {"delta": 0.0, "raw": 0, "median": 0, "population": 4}, {}) == ["match"]
    assert classify("completion", empty, {"delta": 1.0, "raw": 1, "median": 0, "population": 4}, {}) == ["sheet_source_empty"]
    fresh = {"delta": 0, "raw": None, "median": None, "in_cohort": None, "source": "freshness block CA5:CA51 empty"}
    assert classify("freshness", fresh, {"delta": 1.5, "raw": 0, "median": -1.5, "population": 4}, {}) == ["sheet_source_empty"]
    assert classify("freshness", dict(fresh, source=None), {"delta": 1.5, "raw": 0, "median": -1.5, "population": 4}, {}) == ["UNEXPLAINED"]
    cell = lambda sr, ar, classes, n=4: {"sheet": {"raw": sr, "cohort_size": n}, "app": {"raw": ar, "population": 4}, "classes": classes}
    assert median_shift([cell(0, 0, ["UNEXPLAINED"]), cell(9, 4, ["logs_outside_period"]), cell(3, 3, ["match"])]) ==         [["logs_outside_period (via median)"], ["logs_outside_period"], ["match"]]
    assert median_shift([cell(0, 0, ["UNEXPLAINED"], 2), cell(9, 4, ["logs_outside_period"])])[0] == ["UNEXPLAINED"]
    assert median_shift([cell(0, 1, ["UNEXPLAINED"]), cell(9, 4, ["logs_outside_period"])])[0] == ["UNEXPLAINED"]
    assert median_shift([cell(0, 0, ["UNEXPLAINED"]), cell(9, 4, ["UNEXPLAINED"])])[0] == ["UNEXPLAINED"]
    print("PASS reconcile_sheet_scores self-test")


# --------------------------------------------------------------------------- workbook

def load_seed_tool():
    spec = importlib.util.spec_from_file_location("seed_sheet_to_project", ROOT / "tools" / "seed_sheet_to_project.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def read_workbook(seed, path):
    """People and tasks through the seed tool's read_sheet; the two value sheets from the same load."""
    import openpyxl
    original, cache = openpyxl.load_workbook, {}

    def cached(target, **kwargs):  # read_sheet loads data_only=True; reuse that one load
        key = (str(target), tuple(sorted(kwargs.items())))
        if key not in cache:
            cache[key] = original(target, **kwargs)
        return cache[key]
    openpyxl.load_workbook = cached
    try:
        sheet = seed.read_sheet(path)
        workbook = cached(path, data_only=True)
    finally:
        openpyxl.load_workbook = original

    def cells(name, last_row):
        out = {}
        for row in workbook[name].iter_rows(min_row=1, max_row=last_row):
            for c in row:
                if c.value is not None and c.value != "":
                    out[c.coordinate] = c.value
        return out
    tasks = workbook["2_Tasks"]
    task_rows = [{"tag": tasks.cell(r, 1).value, "scale": tasks.cell(r, 5).value, "progress": tasks.cell(r, 6).value,
                  "assignee": (tasks.cell(r, 9).value or "").strip()}
                 for r in range(6, 33) if tasks.cell(r, 1).value]
    return sheet, cells(PERF_SHEET, 24), cells(SUMMARY_SHEET, SUMMARY_LAST_ROW), task_rows


def task_scaled(task_rows, people, two):
    """Per employee: the tasks' scale x progress / 100 as the sheet holds it, and as imported (2-decimal progress)."""
    by_name = {p["name"]: p["employee_id"] for p in people}
    out = {}
    for t in task_rows:
        employee = by_name.get(t["assignee"])
        if employee is None or not is_number(t["scale"]) or not is_number(t["progress"]):
            continue
        sheet, imported = out.setdefault(employee, [Decimal(0), Decimal(0)])
        out[employee] = [sheet + dec(t["scale"]) * dec(t["progress"]) / 100, imported + dec(t["scale"]) * two(t["progress"]) / 100]
    return out


# --------------------------------------------------------------------------- app reads

def read_only(seed, body, result_expr):
    """Like the seed tool's transaction(), but READ ONLY: a stray write fails instead of committing."""
    nonce = uuid.uuid4().hex
    return (f"BEGIN READ ONLY;\nSET LOCAL search_path TO '';\n{body}\n{seed.AS_OWNER}"
            f"SELECT set_config('recon.result', jsonb_build_object('nonce',{seed.lit(nonce)},'value',({result_expr}))::text, false);\n"
            f"COMMIT;\nSELECT current_setting('recon.result')::jsonb AS result;\n-- nonce {nonce}\n")


def app_reads(seed, backend, sheet, as_of):
    """Director id, the RPC's response at as_of and at the default live as-of, and log-credit evidence."""
    lit = seed.lit
    people = sheet["people"]
    director = next(p for p in people if p["role"] == "managing_director")
    ids = backend.query("people", read_only(seed, "", (
        "(SELECT coalesce(jsonb_object_agg(employee_id, jsonb_build_object('id',id,'email',lower(email),'active',is_active,'role',role)),'{}') "
        "FROM public.users WHERE employee_id = ANY(" + lit("{" + ",".join(p["employee_id"] for p in people) + "}") + "::text[]))")))
    missing = [p["employee_id"] for p in people if p["employee_id"] not in ids]
    if missing:
        raise RuntimeError(f"Sheet people missing in the database: {missing}")
    period = dt.date.fromisoformat(sheet["period"])
    nxt = seed.shift_date(period, 1)
    # The live report's default as-of for a past month: its last instant (viewContext), Asia/Kolkata.
    default_as_of = f"({lit(nxt)}::date::timestamp AT TIME ZONE 'Asia/Kolkata') - interval '1 millisecond'"
    chosen = "clock_timestamp() - interval '1 second'" if as_of == "now" else f"{lit(as_of)}::timestamptz"
    body = (seed.as_actor(ids[director["employee_id"]]["id"]) +
            f"SELECT set_config('recon.perf', public.performance_source_context({lit(period)}::date,{lit(nxt)}::date,{chosen})::text, true);\n"
            f"SELECT set_config('recon.default', public.performance_source_context({lit(period)}::date,{lit(nxt)}::date,{default_as_of})::text, true);\n")
    evidence_expr = f"""(WITH latest AS (SELECT DISTINCT ON(source_key) data FROM taskpulse_private.performance_observations
   WHERE kind='log' AND observed_at <= (current_setting('recon.perf')::jsonb->>'asOf')::timestamptz ORDER BY source_key, observed_at DESC, revision DESC),
 credit AS (SELECT (l.data->>'businessDate')::date d, s.key uid,
   CASE WHEN (l.data->>'appliedProgress')::numeric=0 THEN 0 ELSE (l.data->>'appliedProgress')::numeric*(l.data->>'capturedScale')::numeric/100*s.value::numeric/1000000 END amount
   FROM latest l CROSS JOIN LATERAL jsonb_each_text(CASE WHEN jsonb_typeof(l.data->'capturedShares')='object' THEN l.data->'capturedShares' ELSE '{{}}' END) s
   WHERE l.data->>'reversedAt' IS NULL)
 SELECT jsonb_build_object(
  'perf', current_setting('recon.perf')::jsonb,
  'default', current_setting('recon.default')::jsonb,
  'contract_from', (SELECT available_from FROM taskpulse_private.performance_contract),
  'credit', (SELECT coalesce(jsonb_object_agg(uid, x),'{{}}') FROM (SELECT uid, jsonb_build_object(
     'in_period', sum(amount) FILTER (WHERE d >= {lit(period)}::date AND d < {lit(nxt)}::date)::text,
     'after', sum(amount) FILTER (WHERE d >= {lit(nxt)}::date)::text,
     'before', sum(amount) FILTER (WHERE d < {lit(period)}::date)::text) x FROM credit GROUP BY uid) q),
  'logs_by_window', (SELECT jsonb_object_agg(w, n) FROM (SELECT CASE WHEN log_date < {lit(period)}::date THEN 'before'
     WHEN log_date >= {lit(nxt)}::date THEN 'after' ELSE 'in_period' END w, count(*) n FROM public.daily_task_logs
     WHERE reversed_at IS NULL AND log_source='Imported' GROUP BY 1) q),
  'sales_entries', (SELECT count(*) FROM public.sales_logs),
  'attendance_rows', (SELECT count(*) FROM public.attendance_records WHERE attendance_date >= {lit(period)}::date AND attendance_date < {lit(nxt)}::date)
 ))"""
    data = backend.query("performance", read_only(seed, body, evidence_expr))
    return ids, data


def run_engine(sources, workdir):
    src, out = workdir / f"sources-{uuid.uuid4().hex[:6]}.json", workdir / f"calculation-{uuid.uuid4().hex[:6]}.json"
    src.write_text(json.dumps(sources), encoding="utf-8")
    node = "node.exe" if sys.platform == "win32" else "node"
    result = subprocess.run([node, str(ROOT / "tools" / "reconcile_app_scores.mjs"), str(src), str(out)], cwd=ROOT,
                            capture_output=True, text=True, encoding="utf-8")
    if result.returncode:
        raise RuntimeError(f"engine evaluation failed: {result.stderr.strip()[-1500:]}")
    return json.loads(out.read_text(encoding="utf-8"))["calculation"]


# --------------------------------------------------------------------------- comparison

def app_factor(member, factor):
    r = member["factors"][factor]
    return {"raw": r["rawExact"] if r["rawExact"] is not None else r["rawValue"],
            "median": r["medianExact"] if r["medianExact"] is not None else r["median"],
            "delta": r["deltaExact"] if r["deltaExact"] is not None else r["delta"],
            "available": r["available"], "reasons": r["reasons"], "population": len(r["population"])}


def reconcile(sheet_rows, ids, calculation, evidence, scaled):
    members = {m["userId"]: m for m in calculation["members"]}
    result = []
    for employee, entry in sheet_rows.items():
        uid = ids[employee]["id"]
        member = members.get(uid)
        credit = evidence["credit"].get(uid, {})
        outside = sum(dec(credit.get(k) or 0) for k in ("after", "before"))
        sheet_scaled, imported_scaled = scaled.get(employee, [Decimal(0), Decimal(0)])
        member_evidence = {"outside_credit": outside, "in_credit": dec(credit.get("in_period") or 0),
                           "sheet_task_scaled": sheet_scaled, "imported_task_scaled": imported_scaled,
                           "sales_entries": evidence["sales_entries"], "history": evidence["perf"]["historyAvailable"]}
        rows = []
        for f in FACTORS:
            app = app_factor(member, f) if member else {"raw": None, "median": None, "delta": None, "available": False,
                                                         "reasons": ["member_not_in_app_population"], "population": None}
            status, diff = compare(entry["factors"][f]["delta"], app["delta"], FACTOR_PLACES)
            rows.append({"factor": f, "sheet": entry["factors"][f], "app": app, "status": status, "diff": diff,
                         "classes": classify(f, entry["factors"][f], app, member_evidence)})
        composite = None if not member else (member["compositeExact"] if member["compositeExact"] is not None else member["composite"])
        status, diff = compare(entry["score"], composite, SCORE_PLACES)
        score_classes = ["match"] if status == "match" else \
            list(dict.fromkeys(c for r in rows if r["status"] != "match" for c in r["classes"])) or ["UNEXPLAINED"]
        result.append({"employee": employee, "name": entry["name"], "uid": uid, "evidence": member_evidence, "member": member,
                       "attendance_sheet": entry["attendance"], "rows": rows, "composite": composite, "sheet_score": entry["score"]})
    for i, f in enumerate(FACTORS):
        for entry, classes in zip(result, median_shift([e["rows"][i] for e in result])):
            entry["rows"][i]["classes"] = classes
    for entry in result:
        member = entry.pop("member")
        status, diff = compare(entry["sheet_score"], entry["composite"], SCORE_PLACES)
        score_classes = ["match"] if status == "match" else             list(dict.fromkeys(c for r in entry["rows"] if r["status"] != "match" for c in r["classes"])) or ["UNEXPLAINED"]
        entry["score"] = {"sheet": entry.pop("sheet_score"), "app": entry.pop("composite"), "status": status, "diff": diff,
                          "classes": score_classes, "subtotal": member["knownSubtotalExact"] if member else None,
                          "eligible": member["eligibleForRanking"] if member else None,
                          "status_app": member["status"] if member else "absent"}
    return result


def table(entry):
    lines = [f"### {entry['name']} ({entry['employee']})", "",
             "| Factor | Sheet raw | App raw | Sheet median | App median | Sheet diff | App diff | App - sheet | Class | App unavailable reason |",
             "| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | --- | --- |"]
    for r in entry["rows"]:
        s, a = r["sheet"], r["app"]
        reason = ", ".join(a["reasons"]) if not a["available"] else ""
        lines.append(f"| {r['factor']} | {fmt(s['raw'], FACTOR_PLACES)} | {fmt(a['raw'], FACTOR_PLACES)} | {fmt(s['median'], FACTOR_PLACES)} | "
                     f"{fmt(a['median'], FACTOR_PLACES)} | {fmt(s['delta'], FACTOR_PLACES)} | {fmt(a['delta'], FACTOR_PLACES)} | "
                     f"{fmt(r['diff'], FACTOR_PLACES)} | {', '.join(r['classes'])} | {reason} |")
    sc = entry["score"]
    app_score = fmt(sc["app"], SCORE_PLACES) if sc["app"] is not None else f"- (subtotal {fmt(sc['subtotal'], SCORE_PLACES)})"
    lines.append(f"| **score** | | | | | {fmt(sc['sheet'], SCORE_PLACES)} | {app_score} | {fmt(sc['diff'], SCORE_PLACES)} | "
                 f"{', '.join(sc['classes'])} | {'app ' + sc['status_app'] + ('; not eligible for ranking' if sc['eligible'] is False else '')} |")
    return "\n".join(lines)


def print_console(results):
    for entry in results:
        print(f"\n{entry['name']} ({entry['employee']})")
        print(f"  {'factor':<12} {'sheet':>16} {'app':>16} {'app-sheet':>14}  class / app reason")
        for r in entry["rows"]:
            reason = "" if r["app"]["available"] else "  [" + ", ".join(r["app"]["reasons"]) + "]"
            print(f"  {r['factor']:<12} {fmt(r['sheet']['delta'], FACTOR_PLACES):>16} {fmt(r['app']['delta'], FACTOR_PLACES):>16} "
                  f"{fmt(r['diff'], FACTOR_PLACES):>14}  {', '.join(r['classes'])}{reason}")
        sc = entry["score"]
        print(f"  {'SCORE':<12} {fmt(sc['sheet'], SCORE_PLACES):>16} {fmt(sc['app'], SCORE_PLACES):>16} {fmt(sc['diff'], SCORE_PLACES):>14}  "
              f"{', '.join(sc['classes'])}  (app {sc['status_app']}, available-factor subtotal {fmt(sc['subtotal'], SCORE_PLACES)})")


def report(path, args, sheet, sheet_sha, results, evidence, default_calc, target_label):
    perf = evidence["perf"]
    counts = {}
    for entry in results:
        for r in entry["rows"]:
            for c in r["classes"]:
                c = c.replace(" (via median)", "")
                counts[c] = counts.get(c, 0) + 1
    unexplained = [(e, r) for e in results for r in e["rows"] if "UNEXPLAINED" in r["classes"]]
    default_members = len(default_calc["members"]) if default_calc else 0
    lines = [
        "# M17 reconciliation: September 2026 workbook vs app", "",
        f"Generated {dt.datetime.now(dt.timezone.utc):%Y-%m-%d %H:%M} UTC by `tools/reconcile_sheet_scores.py` "
        f"against {target_label}. Scoring rules were not changed; every difference below is classified.", "",
        "## What was compared", "",
        f"- Workbook: sha256 `{sheet_sha[:16]}…`, period {sheet['period'][:7]}, {len(sheet['people'])} people, "
        f"{len(sheet['tasks'])} tasks, {sum(len(t['logs']) for t in sheet['tasks'])} logs. Per-employee values are "
        f"`{PERF_SHEET}` rows 3-{2 + len(results)} (differences with median, L = Sum); raw values and medians come from "
        f"`{SUMMARY_SHEET}` (rows tied to a person by an exactly matching difference, because the export drops the "
        "dynamic-array name columns).",
        f"- App: `public.performance_source_context({sheet['period']}, next month, as-of)` read as the director in a READ ONLY "
        "transaction, evaluated by `parsePerformanceSources` + `resolvePerformance` from `modules/scoring/client.ts` "
        "(what `loadPerformance` does after its RPC).",
        f"- As-of used for the table: `{perf['asOf']}` (local date {perf['asOfLocalDate']}); observation history available "
        f"from `{evidence['contract_from']}`; history available: {perf['historyAvailable']}.",
        f"- Live report default as-of for September (the month's last instant): history available "
        f"{evidence['default']['historyAvailable']}, {default_members} members in the calculation. "
        + ("The app has no observations from before its history baseline, so a September view at its default as-of "
           "shows no September values: every cell of that view is an `as_of_window` difference. The tables "
           "below therefore use the later as-of above, after the seed's observations."
           if not evidence["default"]["historyAvailable"] else ""),
        f"- Comparison precision: factors {FACTOR_PLACES} decimals, scores {SCORE_PLACES}, half away from zero. "
        "\"App - sheet\" is the difference of the rounded differences-with-median.",
        f"- Imported logs by date window: {json.dumps(evidence['logs_by_window'], sort_keys=True)} "
        f"(known: {KNOWN_OUTSIDE_LOGS['october']} October-dated, {KNOWN_OUTSIDE_LOGS['2006']} dated 2006). "
        f"September attendance rows: {evidence['attendance_rows']}; sales_logs rows: {evidence['sales_entries']}.", "",
        "## Summary", ""]
    for f in FACTORS:
        cells = [e["rows"][FACTORS.index(f)] for e in results]
        equal_raw = sum(compare(c["sheet"].get("raw"), c["app"].get("raw"), FACTOR_PLACES)[0] == "match" for c in cells)
        known_raw = sum(c["sheet"].get("raw") is not None and c["app"].get("raw") is not None for c in cells)
        matches = sum(c["status"] == "match" for c in cells)
        classes = sorted({k.replace(" (via median)", "") for c in cells for k in c["classes"]} - {"match"})
        lines.append(f"- {f}: {matches}/{len(cells)} differences equal; raw values equal for {equal_raw}/{known_raw} members "
                     f"where both sides have one; causes: {', '.join(classes) or 'none'}.")
    incomplete = [e["name"] for e in results if e["score"]["app"] is None]
    lines += [f"- Score: the app's composite is unavailable for {len(incomplete)}/{len(results)} people "
              f"({', '.join(incomplete) or 'none'}): productivity needs September attendance, so nobody is complete or "
              "eligible for ranking. The available-factor subtotal is shown for information only; it is not a score.",
              "- The workbook's own L Sum equals its 1_2 factors with the workbook weights (checked for every row).", "",
              "## Tables", ""]
    for entry in results:
        ev = entry["evidence"]
        lines += [table(entry), "",
                  f"Evidence: September-dated log credit {fmt(ev['in_credit'], FACTOR_PLACES)}, credit on logs dated outside "
                  f"September {fmt(ev['outside_credit'], FACTOR_PLACES)} (scaled progress units); the member's tasks' scale x "
                  f"progress {fmt(ev['sheet_task_scaled'], FACTOR_PLACES)} in the sheet, {fmt(ev['imported_task_scaled'], FACTOR_PLACES)} "
                  f"as imported (2-decimal progress); sheet attendance hours {fmt(entry['attendance_sheet'], 2)}.", ""]
    lines += ["## Classification", "", "| Class | Cells | Meaning |", "| --- | ---: | --- |"]
    for c in ["match", *CLASSES]:
        if counts.get(c):
            lines.append(f"| {c} | {counts[c]} | {CLASSES.get(c, 'Equal at display precision.')} |")
    lines += ["", "## UNEXPLAINED items (questions for the user)", ""]
    if unexplained:
        for e, r in unexplained:
            lines.append(f"- {e['name']} {r['factor']}: sheet {fmt(r['sheet']['delta'], FACTOR_PLACES)}, app "
                         f"{fmt(r['app']['delta'], FACTOR_PLACES)} (raw sheet {fmt(r['sheet']['raw'], FACTOR_PLACES)} / app "
                         f"{fmt(r['app']['raw'], FACTOR_PLACES)}). Which value is right, and what input explains it?")
    else:
        lines.append("- None: every difference has a known cause.")
    lines += ["", "## Inputs and decisions still open (not differences in the engine)", "",
              "- September attendance is not entered, so productivity and every composite score stay unavailable. "
              "Will you enter September attendance (or supply the attendance workbook the sheet imports)?",
              "- Progress: the seed credits each task's whole sheet progress on its latest-dated log, so tasks whose last "
              "log is October-dated earn no September credit. The sheet has no per-log progress history. Should the "
              "September import instead credit progress on the last September-dated log, or is the October dating correct?",
              "- The 5 logs dated 2006-09-24 on Sep26_T6 look like a year typo for 2026-09-24. Confirm before correcting them.",
              "- The sheet's freshness block (CA5:CA51) and completion/facts/ideas/suggestions/sales blocks are empty in this "
              "export (medians #NUM!), so the workbook adds 0 for those factors; the app measures freshness from task dates "
              "as of the as-of day. Is the empty freshness block expected for September?"]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# --------------------------------------------------------------------------- local build

def build_local(seed, args, workdir, sheet):
    """Fresh disposable database from a migrations folder, then the seed tool applies the sheet.

    As in tools/test_sheet_seed_database.py, the sheet's director exists first (the seed tool
    impersonates an existing active director); the seed tool creates the other people."""
    sys.path.insert(0, str(ROOT / "tools"))
    import ast
    source = (ROOT / "tools" / "test_m09_database.py").read_text(encoding="utf-8")
    bootstrap = next(ast.literal_eval(node.value) for node in ast.parse(source).body
                     if isinstance(node, ast.Assign) and getattr(node.targets[0], "id", "") == "bootstrap")
    migrations = sorted(args.migrations.glob("*.sql"), key=lambda p: int(p.name.split("_")[0]))

    def psql(db, statement):
        result = subprocess.run([args.psql, "-X", "-q", "-h", "127.0.0.1", "-p", str(args.port), "-U", "postgres", "-d", db,
                                 "-v", "ON_ERROR_STOP=1", "-At"], input=statement, text=True, capture_output=True, encoding="utf-8")
        if result.returncode:
            raise RuntimeError(result.stderr.strip()[-1500:])
    psql("postgres", f"DROP DATABASE IF EXISTS {args.db} WITH (FORCE);")
    psql("postgres", f"CREATE DATABASE {args.db};")
    psql(args.db, bootstrap)
    for migration in migrations:
        psql(args.db, migration.read_text(encoding="utf-8"))
    psql(args.db, "CREATE SCHEMA supabase_migrations; CREATE TABLE supabase_migrations.schema_migrations(version text PRIMARY KEY, name text);\n"
         + "".join(f"INSERT INTO supabase_migrations.schema_migrations VALUES('{p.name.split('_')[0]}','{p.stem}');\n" for p in migrations))
    director = next(p for p in sheet["people"] if p["role"] == "managing_director")
    meta = json.dumps({"employee_id": director["employee_id"], "full_name": director["name"]})
    psql(args.db, f"INSERT INTO auth.users(id,email,raw_user_meta_data) VALUES(gen_random_uuid(),{seed.lit(director['email'])},{seed.lit(meta)}::jsonb);\n"
         f"UPDATE public.users SET role='managing_director' WHERE lower(email)={seed.lit(director['email'])};")
    print(f"Built {args.db}:{migrations[0].name}..{migrations[-1].name} ({len(migrations)} migrations)")
    command = [sys.executable, "-I", str(ROOT / "tools" / "seed_sheet_to_project.py"), "--xlsx", str(args.xlsx), "--target", "local",
               "--psql", args.psql, "--port", str(args.port), "--db", args.db, "--allow-ledger", str(len(migrations)),
               "--manifest-dir", str(workdir), "--yes"]
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
    if result.returncode or "PASS Phase D verification" not in result.stdout:
        raise RuntimeError(f"seed failed (exit {result.returncode}):\n{result.stdout[-3000:]}{result.stderr[-1500:]}")
    print("Seeded the sheet: PASS Phase D verification")


# --------------------------------------------------------------------------- main

def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--xlsx", type=Path)
    parser.add_argument("--target", choices=("local", "hosted"), default="local")
    parser.add_argument("--psql")
    parser.add_argument("--port", type=int, default=15432)
    parser.add_argument("--db", default="taskpulse_reconcile_test")
    parser.add_argument("--build", action="store_true", help="local only: recreate --db from --migrations and seed the sheet")
    parser.add_argument("--migrations", type=Path)
    parser.add_argument("--project-ref")
    parser.add_argument("--workdir", type=Path, help="hosted: Supabase CLI folder linked to --project-ref")
    parser.add_argument("--as-of", default="now", help="'now' (default) or an ISO timestamp")
    parser.add_argument("--report", type=Path, help="write the Markdown report here")
    parser.add_argument("--out-dir", type=Path, default=Path(tempfile.gettempdir()) / "taskpulse-reconcile")
    args = parser.parse_args(argv)
    if args.self_test:
        return args
    if not args.xlsx:
        parser.error("--xlsx is required")
    if args.target == "local":
        if not args.psql or not re.fullmatch(r"taskpulse_[a-z0-9_]+_test", args.db) or not 1024 < args.port < 65536:
            parser.error("--target local needs --psql and a disposable --db named taskpulse_*_test")
        if args.build and not args.migrations:
            parser.error("--build needs --migrations")
    elif args.build or not (args.project_ref and args.workdir):
        parser.error("--target hosted needs --project-ref and --workdir and never --build")
    out_dir = args.out_dir.resolve()
    if out_dir == ROOT or ROOT in out_dir.parents:
        parser.error("--out-dir must be outside the repository")
    args.out_dir = out_dir
    return args


def main(argv=None):
    args = parse_args(argv)
    if args.self_test:
        self_test()
        return 0
    self_test()
    seed = load_seed_tool()
    args.out_dir.mkdir(parents=True, exist_ok=True)
    workdir = Path(tempfile.mkdtemp(prefix="run-", dir=args.out_dir))
    sheet, perf_cells, summary_cells, task_rows = read_workbook(seed, args.xlsx)
    if args.build:
        build_local(seed, args, workdir, sheet)
    backend = seed.HostedBackend(workdir, args.project_ref, args.workdir) if args.target == "hosted" else \
        seed.LocalBackend(workdir, args.psql, args.port, args.db)
    sheet_rows = sheet_values(perf_cells, summary_cells, sheet["people"])
    for employee, entry in sheet_rows.items():
        check = sheet_score_check(entry)
        if rounded(check, SCORE_PLACES) != rounded(entry["score"], SCORE_PLACES):
            raise RuntimeError(f"{entry['name']}: the workbook's L Sum {entry['score']} does not equal its own factors ({check})")
    ids, evidence = app_reads(seed, backend, sheet, args.as_of)
    calculation = run_engine(evidence["perf"], workdir)
    default_calc = run_engine(evidence["default"], workdir) if evidence["default"]["members"] else None
    results = reconcile(sheet_rows, ids, calculation, evidence, task_scaled(task_rows, sheet["people"], seed.two))
    print(f"Target: {backend.label}; as-of {evidence['perf']['asOf']}; history available {evidence['perf']['historyAvailable']}")
    print_console(results)
    if args.report:
        report(args.report, args, sheet, seed.sha256_file(args.xlsx), results, evidence, default_calc, backend.label)
        print(f"\nReport: {args.report}")
    (workdir / "results.json").write_text(json.dumps(results, indent=1, default=str), encoding="utf-8")
    unexplained = sum("UNEXPLAINED" in r["classes"] for e in results for r in e["rows"])
    print(f"\n{unexplained} UNEXPLAINED cell(s); details in {workdir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
