"""Seed a TaskPulse project with the September 2026 Google Sheets tracker as test data.

Reads the tracker workbook (an .xlsx file, or the link-shared export downloaded
into its own empty directory), then plans and, only with --yes, applies:

  A  people   confirmed Auth users without password/invite (trigger-created profiles, roles set as owner)
  B  replace  delete the named existing tasks and every dependent row (snapshotted to the manifest first)
  C  tasks    one transaction per sheet task through the app's own RPCs, impersonating the director/assignee
  D  verify   counts, hours, eligibility, triggers and the director's performance/presentation reads

Without --yes the tool only prints the plan (dry run). Hosted mode talks to the
project through the Supabase CLI query channel and, for Auth users, the Admin
API with a service-role key fetched into memory (never printed or written).
Local mode targets a disposable loopback PostgreSQL with the JWT shim of
tools/test_m*_database.py. Every write is recorded in a JSON manifest outside
the repository; --undo MANIFEST removes the rows that run created.
"""

import argparse
import datetime as dt
import hashlib
import json
import re
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
import uuid
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PRODUCTION_REF = "sbtxtsisjhuciobxykcm"
DEFAULT_LEDGER = 46
NS = uuid.UUID("5e3d7f0a-2c41-4b8e-9a51-0d6c1f9e2b77")  # fixed namespace for this tool's operation ids
TWO = Decimal("0.01")

ROLE_MAP = {"Managing Director": "managing_director", "Proj. Engineer": "project_engineer",
            "Data Manager": "data_manager", "Team Member": "team_member"}
STATUSES = {"Not Started", "In Progress", "Discarded"}
PRIORITIES = {"Low", "Medium", "High", "Focus", "No"}
# Logs per task in the September 2026 sheet; any drift stops the run.
EXPECTED_LOGS = {"T6": 128, "T5": 80, "AD22": 61, "PJT27": 57, "AD21": 23, "MKT20": 9, "T25": 9,
                 "T8": 8, "PJT15": 3, "T26": 3, "T9": 1}
EXPECTED_TOTAL_LOGS = 382
EXPECTED_VALID = {"T5", "T6", "T8", "MKT20", "AD21", "AD22"}
DISCARD_REASON = "Discarded in the September 2026 sheet tracker (imported test data)"
REPLACE_REASON = "Replaced by the September 2026 sheet seed (test data)"
UNDO_REASON = "Removed by sheet-seed undo"
# Rows in these tables belong to people, not to the task; deleting a task must never take them.
FOREIGN_OWNED = ("public.daily_facts", "public.ideas", "public.suggestions", "public.sales_logs", "public.meeting_points")
GUARDED_TRIGGERS = (("public.daily_task_logs", "m02_retain_logs"), ("public.tasks", "m02_task_guard"),
                    ("public.monthly_performance_snapshots", "m07_retain_month"),
                    ("public.report_snapshots", "m07_retain_report"), ("public.task_timers", "m02_retain_timers"))
SECRET_PATTERN = re.compile(r"(eyJ[\w-]{10,}\.[\w-]+\.[\w-]+|sb_(secret|publishable)_\w+|postgres(ql)?://\S+|sbp_\w+)")


class Refusal(Exception):
    """A gate failed; nothing (further) is written."""


# --------------------------------------------------------------------------- helpers

def lit(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def jlit(value):
    return lit(json.dumps(value, ensure_ascii=False, sort_keys=True)) + "::jsonb"


def uuid_array(ids):
    return "ARRAY[" + ",".join(lit(i) for i in ids) + "]::uuid[]" if ids else "ARRAY[]::uuid[]"


def op_id(*parts):
    return str(uuid.uuid5(NS, "|".join(str(p) for p in parts)))


def dec(value):
    return Decimal(repr(value)) if isinstance(value, float) else Decimal(str(value))


def two(value):
    return dec(value).quantize(TWO, rounding=ROUND_HALF_UP)


def fmt_num(value):
    text = format(dec(value).normalize(), "f")
    return text


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1 << 16), b""):
            digest.update(block)
    return digest.hexdigest()


def redact(text):
    return SECRET_PATTERN.sub("[redacted]", text or "")


def shift_date(day, months):
    if not months:
        return day
    index = day.year * 12 + day.month - 1 + months
    year, month = divmod(index, 12)
    month += 1
    last = (dt.date(year + (month == 12), month % 12 + 1, 1) - dt.timedelta(days=1)).day
    return dt.date(year, month, min(day.day, last))


def as_date(value, what):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    raise Refusal(f"{what}: expected a date, found {value!r}")


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def as_actor(uid):
    claims = json.dumps({"sub": uid, "role": "authenticated"})
    return (f"SET LOCAL ROLE authenticated;\nSELECT set_config('request.jwt.claim.sub',{lit(uid)},true),"
            f" set_config('request.jwt.claims',{lit(claims)},true);\n")


AS_OWNER = ("RESET ROLE;\nSELECT set_config('request.jwt.claim.sub','',true), set_config('request.jwt.claims','',true);\n")


TRIGGERS_EXPR = ("(SELECT jsonb_object_agg(n.nspname||'.'||c.relname||'.'||g.tgname, g.tgenabled) FROM pg_trigger g "
                 "JOIN pg_class c ON c.oid=g.tgrelid JOIN pg_namespace n ON n.oid=c.relnamespace WHERE NOT g.tgisinternal AND ("
                 + " OR ".join(f"(n.nspname||'.'||c.relname={lit(r)} AND g.tgname={lit(t)})" for r, t in GUARDED_TRIGGERS) + "))")


def transaction(body, result_expr):
    """One atomic script. search_path is empty so regclass text is always schema-qualified.
    The result is carried across COMMIT in a session setting, tagged with a fresh nonce, so the
    final statement (what the CLI query channel returns) is one jsonb row that cannot be a
    stale value left by an earlier request on the same connection."""
    nonce = uuid.uuid4().hex
    return (f"BEGIN;\nSET LOCAL search_path TO '';\n{body}\n{AS_OWNER}"
            f"SELECT set_config('seed.result', jsonb_build_object('nonce',{lit(nonce)},'value',({result_expr}))::text, false);\n"
            f"COMMIT;\nSELECT current_setting('seed.result')::jsonb AS result;\n-- nonce {nonce}\n")


def unwrap(name, sql, value):
    nonce = re.search(r"-- nonce ([0-9a-f]{32})\n$", sql).group(1)
    if not isinstance(value, dict) or value.get("nonce") != nonce:
        raise RuntimeError(f"{name}: result does not carry this request's nonce")
    return value["value"]


# --------------------------------------------------------------------------- backends

class Backend:
    label = ""
    is_production = False

    def __init__(self, sql_dir):
        self.sql_dir = sql_dir
        self.counter = 0

    def _file(self, name, sql):
        self.counter += 1
        path = self.sql_dir / f"{self.counter:03d}-{name}.sql"
        path.write_text(sql, encoding="utf-8")
        return path

    def query(self, name, sql):
        raise NotImplementedError


class HostedBackend(Backend):
    def __init__(self, sql_dir, ref, workdir):
        super().__init__(sql_dir)
        self.ref, self.workdir = ref, workdir
        self.label = f"hosted project {ref}"
        self.is_production = ref == PRODUCTION_REF
        self._service_key = None

    def cli(self, *arguments):
        executable = "supabase.cmd" if sys.platform == "win32" else "supabase"
        result = subprocess.run([executable, *arguments], capture_output=True, text=True, encoding="utf-8", check=False)
        if result.returncode:
            raise RuntimeError(f"Supabase CLI {arguments[0]} {arguments[1] if len(arguments) > 1 else ''} failed "
                               f"(exit {result.returncode}): {redact(result.stderr.strip())[-1500:]}")
        output = result.stdout
        first = min((p for p in (output.find("{"), output.find("[")) if p >= 0), default=-1)
        if first < 0:
            raise RuntimeError(f"Supabase CLI returned no JSON for {arguments[0]}")
        return json.loads(output[first:])

    def query(self, name, sql):
        path = self._file(name, sql)
        data = self.cli("db", "query", "--linked", "--project-ref", self.ref, "--workdir", str(self.workdir),
                        "--file", str(path))
        rows = data.get("rows") if isinstance(data, dict) else data
        if not rows or "result" not in rows[-1]:
            raise RuntimeError(f"{name}: query returned no result row")
        value = rows[-1]["result"]
        return unwrap(name, sql, json.loads(value) if isinstance(value, str) else value)

    def projects(self):
        return self.cli("projects", "list", "--output", "json")

    def _key(self):
        if self._service_key is None:
            keys = self.cli("projects", "api-keys", "--project-ref", self.ref, "--reveal", "--output", "json")
            self._service_key = next((k["api_key"] for k in keys if k.get("name") == "service_role"), None)
            if not self._service_key:
                raise RuntimeError("The project's service_role key is unavailable")
        return self._service_key

    def request(self, method, path, payload=None):
        key = self._key()
        body = json.dumps(payload).encode() if payload is not None else None
        headers = {"apikey": key, "Authorization": f"Bearer {key}", "Accept": "application/json"}
        if body is not None:
            headers["Content-Type"] = "application/json"
        req = urllib.request.Request(f"https://{self.ref}.supabase.co{path}", data=body, headers=headers, method=method)
        try:
            with urllib.request.urlopen(req, timeout=30) as response:
                content = response.read()
                return response.status, json.loads(content) if content else None
        except urllib.error.HTTPError as error:
            # Keep only a machine-readable code; bodies can echo inputs.
            try:
                code = json.loads(error.read()).get("error_code")
            except (ValueError, AttributeError):
                code = None
            return error.code, {"code": code}

    def create_auth_user(self, email, user_meta, app_meta):
        status, data = self.request("POST", "/auth/v1/admin/users", {
            "email": email, "email_confirm": True, "user_metadata": user_meta, "app_metadata": app_meta})
        if status not in (200, 201) or not isinstance(data, dict) or not data.get("id"):
            raise RuntimeError(f"Auth user creation failed for {email}: HTTP {status} {data}")
        return data["id"]

    def delete_auth_user(self, uid):
        status, data = self.request("DELETE", f"/auth/v1/admin/users/{uid}")
        if status not in (200, 204):
            raise RuntimeError(f"Auth user deletion failed for {uid}: HTTP {status} {data}")


class LocalBackend(Backend):
    def __init__(self, sql_dir, psql, port, db):
        super().__init__(sql_dir)
        self.psql, self.port, self.db = psql, port, db
        self.label = f"local database {db} on 127.0.0.1:{port}"

    def query(self, name, sql):
        path = self._file(name, sql)
        result = subprocess.run([self.psql, "-X", "-q", "-At", "-h", "127.0.0.1", "-p", str(self.port), "-U", "postgres",
                                 "-d", self.db, "-v", "ON_ERROR_STOP=1", "-f", str(path)],
                                capture_output=True, text=True, encoding="utf-8", check=False)
        if result.returncode:
            raise RuntimeError(f"{name} failed: {result.stderr.strip()[-1500:]}")
        lines = [line for line in result.stdout.splitlines() if line.strip()]
        if not lines:
            raise RuntimeError(f"{name}: query returned no result row")
        return unwrap(name, sql, json.loads(lines[-1]))

    def create_auth_user(self, email, user_meta, app_meta):
        # The loopback shim's auth.users stands in for GoTrue; its trigger creates the profile.
        return self.query("auth-user", transaction(
            "INSERT INTO auth.users(id,email,raw_user_meta_data,raw_app_meta_data) VALUES(gen_random_uuid(),"
            f"{lit(email)},{jlit(user_meta)},{jlit(app_meta)});", f"(SELECT to_jsonb(id) FROM auth.users WHERE email={lit(email)})"))

    def delete_auth_user(self, uid):
        self.query("auth-delete", transaction(f"DELETE FROM auth.users WHERE id={lit(uid)}::uuid;", "to_jsonb(true)"))


# --------------------------------------------------------------------------- sheet

def load_categories():
    text = (ROOT / "shared" / "tracker" / "enums.ts").read_text(encoding="utf-8")
    block = re.search(r"TASK_CATEGORIES = \[(.*?)\]", text, re.S)
    if not block:
        raise Refusal("shared/tracker/enums.ts: TASK_CATEGORIES not found")
    return set(re.findall(r"'([A-Z]+)'", block.group(1)))


def read_sheet(path, email_domain=None, shift_months=0):
    import warnings
    import openpyxl
    warnings.filterwarnings("ignore", module="openpyxl")
    workbook = openpyxl.load_workbook(path, data_only=True)
    warnings_out = []
    for name in ("C_Team Members", "A_Sheet_Attributes", "2_Tasks"):
        if name not in workbook.sheetnames:
            raise Refusal(f"Workbook has no '{name}' sheet")

    attributes = workbook["A_Sheet_Attributes"]
    month_name, year = attributes["B1"].value, attributes["B2"].value
    prefix = attributes["C3"].value
    try:
        period = dt.datetime.strptime(f"{month_name} {int(year)}", "%B %Y").date()
    except (TypeError, ValueError):
        raise Refusal(f"A_Sheet_Attributes B1:B2 do not name a month and year: {month_name!r} {year!r}")
    if prefix != period.strftime("%b%y"):
        raise Refusal(f"A_Sheet_Attributes C3 tag prefix {prefix!r} does not match {period:%B %Y}")
    categories = load_categories()
    type_prefix = {}
    for label, code in attributes.iter_rows(min_row=2, min_col=9, max_col=10, values_only=True):
        if blank(label):
            continue
        if code not in categories:
            raise Refusal(f"Task type {label!r} maps to {code!r}, which is not a task_category in shared/tracker/enums.ts")
        type_prefix[label.strip()] = code

    people = []
    members = workbook["C_Team Members"]
    header = [members.cell(1, c).value for c in (1, 2, 5, 6)]
    if header != ["Team Member", "Mail-ID", "Employee ID", "Role"]:
        raise Refusal(f"C_Team Members header differs: {header}")
    for row in members.iter_rows(min_row=2, values_only=True):
        if blank(row[0]):
            continue
        name, email, employee, role = row[0].strip(), (row[1] or "").strip().lower(), (row[4] or "").strip(), row[5]
        if role not in ROLE_MAP or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[a-z]+", email) or not employee or len(employee) > 20:
            raise Refusal(f"C_Team Members row for {name!r} has an unusable email/employee id/role")
        if email_domain:
            email = email.split("@")[0] + "@" + email_domain
        people.append({"name": name, "email": email, "employee_id": employee, "sheet_role": role, "role": ROLE_MAP[role]})
    by_name = {p["name"]: p for p in people}
    if len(by_name) != len(people) or len({p["email"] for p in people}) != len(people) or \
       len({p["employee_id"] for p in people}) != len(people):
        raise Refusal("C_Team Members has duplicate names, emails or employee ids")
    directors = [p for p in people if p["role"] == "managing_director"]
    if len(directors) != 1:
        raise Refusal("C_Team Members must name exactly one Managing Director")

    tasks_sheet = workbook["2_Tasks"]
    expected_header = ["Task Tag", "Task Type", "Task", "Comments", "Task Scale", "Progress %", "Scaled Progress",
                       "Priority", "Assignee", "Reviewer(s)", "Task Productivity", "Task Freshness", "Status",
                       "Task Assigned", "Due on", "Task Done on"]
    header = [tasks_sheet.cell(4, c).value for c in range(1, 17)]
    if header != expected_header:
        raise Refusal(f"2_Tasks header row 4 differs: {header}")
    tasks = []
    for number, row in enumerate(tasks_sheet.iter_rows(min_row=6, values_only=True), start=6):
        if blank(row[0]):
            continue
        if number > 32:
            raise Refusal(f"2_Tasks row {number} has a tag outside data rows 6-32")
        tag = str(row[0]).strip()
        match = re.fullmatch(re.escape(prefix) + r"_([A-Z]+)(\d+)", tag)
        if not match:
            raise Refusal(f"2_Tasks row {number}: tag {tag!r} is not {prefix}_<PREFIX><n>")
        type_label = (row[1] or "").strip()
        category = type_prefix.get(type_label)
        if category is None or category != match.group(1):
            raise Refusal(f"{tag}: type {type_label!r} does not map to the tag prefix {match.group(1)!r}")
        if blank(row[2]):
            raise Refusal(f"{tag}: blank title")
        if not blank(row[15]):
            raise Refusal(f"{tag}: 'Task Done on' must be empty (Done tasks are not imported)")
        status, priority = row[12], row[7]
        if status not in STATUSES or priority not in PRIORITIES:
            raise Refusal(f"{tag}: unsupported status {status!r} or priority {priority!r}")
        assignee = by_name.get((row[8] or "").strip())
        reviewer_names = [n.strip() for n in re.split(r"[,;&/]", row[9] or "") if n.strip()]
        reviewers = [by_name.get(n) for n in reviewer_names]
        if assignee is None or not reviewers or None in reviewers or len(set(reviewer_names)) != len(reviewer_names):
            raise Refusal(f"{tag}: assignee/reviewers are not team members: {row[8]!r} / {row[9]!r}")
        scale, progress = row[4], row[5]
        if not isinstance(scale, (int, float)) or scale < 0 or not isinstance(progress, (int, float)) or not 0 <= progress <= 100:
            raise Refusal(f"{tag}: scale/progress must be numbers (scale >=0, progress 0-100)")
        assigned = shift_date(as_date(row[13], f"{tag} assigned"), shift_months)
        due = shift_date(as_date(row[14], f"{tag} due"), shift_months)
        tasks.append({"row": number, "tag": tag, "short": match.group(1) + match.group(2), "category": category,
                      "title": row[2], "comments": None if blank(row[3]) else str(row[3]),
                      "scale": fmt_num(scale), "progress": str(two(progress)), "priority": priority,
                      "assignee": assignee["email"], "reviewers": [r["email"] for r in reviewers],
                      "status": status, "assigned_date": assigned.isoformat(), "due_date": due.isoformat(),
                      "logs": read_task_page(workbook, tag, shift_months, warnings_out)})
    if len({t["tag"] for t in tasks}) != len(tasks):
        raise Refusal("2_Tasks has duplicate tags")
    counts = {t["short"]: len(t["logs"]) for t in tasks if t["logs"]}
    if counts != EXPECTED_LOGS or sum(counts.values()) != EXPECTED_TOTAL_LOGS:
        raise Refusal(f"Task-log counts differ from the expected September sheet: {counts}")
    return {"period": period.isoformat(), "shifted_period": shift_date(period, shift_months).isoformat(),
            "prefix": prefix, "people": people, "tasks": tasks, "warnings": warnings_out}


def read_task_page(workbook, tag, shift_months, warnings_out):
    if tag not in workbook.sheetnames:
        raise Refusal(f"Task page {tag} is missing")
    page = workbook[tag]
    if page["A1"].value != "Task Tag" or page["B1"].value != tag:
        raise Refusal(f"{tag}: A1/B1 must read 'Task Tag' / '{tag}'")
    total = None
    for r in range(1, 13):
        for c in range(1, 9):
            if page.cell(r, c).value == "Total Hours Worked":
                total = page.cell(r + 1, c).value
    if not isinstance(total, (int, float)):
        raise Refusal(f"{tag}: 'Total Hours Worked' total not found")
    start = next((r for r in range(1, page.max_row + 1) if page.cell(r, 1).value == "TASK LOG"), None)
    if start is None:
        raise Refusal(f"{tag}: no 'TASK LOG' row")
    header = [page.cell(start + 1, c).value for c in (1, 2, 6, 7)]
    if header != ["Date", "Update", "Hours Worked", "Raw Links"]:
        raise Refusal(f"{tag}: log header differs: {header}")
    logs, raw_total = [], Decimal(0)
    for r in range(start + 2, page.max_row + 1):
        day, update, hours, links = (page.cell(r, c).value for c in (1, 2, 6, 7))
        if all(blank(v) for v in (day, update, hours, links)):
            continue
        day = shift_date(as_date(day, f"{tag} row {r}"), shift_months)
        if isinstance(hours, dt.time):
            hours = Decimal(hours.hour) + Decimal(hours.minute) / 60 + Decimal(hours.second) / 3600
        elif isinstance(hours, dt.timedelta):
            hours = Decimal(hours.total_seconds()) / 3600
        elif isinstance(hours, (int, float)):
            hours = dec(hours)
        else:
            raise Refusal(f"{tag} row {r}: hours {hours!r} is not a number/time")
        raw_total += hours
        rounded = hours.quantize(TWO, rounding=ROUND_HALF_UP)
        if rounded != hours:
            warnings_out.append(f"{tag} row {r}: hours {hours} rounded to {rounded}")
        if not Decimal(0) <= rounded <= Decimal(24):
            raise Refusal(f"{tag} row {r}: hours {rounded} outside 0-24")
        notes = None if blank(update) else str(update)
        raw_links, extra = [], []
        for token in ([] if blank(links) else str(links).split()):
            (raw_links if re.fullmatch(r"https?://\S+", token) and len(raw_links) < 10 else extra).append(token)
        if extra:
            notes = ((notes + "\n") if notes else "") + "Links: " + " ".join(extra)
            warnings_out.append(f"{tag} row {r}: {len(extra)} link value(s) kept in notes")
        logs.append({"row": r, "log_date": day.isoformat(), "hours": str(rounded), "notes": notes, "raw_links": raw_links})
    if abs(raw_total - dec(total)) > Decimal("0.005"):
        raise Refusal(f"{tag}: log hours {raw_total} differ from the page total {total}")
    return logs


def log_digest(rows):
    # create_task_log stores btrim(notes) (spaces only), so both sides are compared that way.
    lines = sorted(f"{r['log_date']}|{Decimal(r['hours']):.2f}|{(r['notes'] or '').strip(' ')}|{' '.join(r['raw_links'] or [])}"
                   for r in rows)
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


# --------------------------------------------------------------------------- database reads

def state_sql(sheet):
    emails = [p["email"] for p in sheet["people"]]
    months = sorted({t_log["log_date"][:7] + "-01" for t in sheet["tasks"] for t_log in t["logs"]} |
                    {sheet["shifted_period"]})
    return transaction("", f"""jsonb_build_object(
 'current_user', current_user,
 'definer_member', (SELECT pg_has_role(current_user,'taskpulse_definer','MEMBER') FROM pg_roles WHERE rolname='taskpulse_definer'),
 'ledger', CASE WHEN to_regclass('supabase_migrations.schema_migrations') IS NULL THEN NULL
   ELSE (SELECT count(*) FROM supabase_migrations.schema_migrations) END,
 'organizations', (SELECT coalesce(jsonb_agg(jsonb_build_object('id',id,'name',name,'status',status,'is_default',is_default)),'[]') FROM public.organizations),
 'users', (SELECT coalesce(jsonb_agg(jsonb_build_object('id',id,'email',lower(email),'employee_id',employee_id,'full_name',full_name,
   'role',role,'is_active',is_active,'organization_id',organization_id) ORDER BY email),'[]') FROM public.users),
 'auth_emails', (SELECT coalesce(jsonb_object_agg(lower(email),id),'{{}}') FROM auth.users WHERE lower(email) = ANY({lit('{' + ','.join(emails) + '}')}::text[])),
 'tasks', (SELECT coalesce(jsonb_agg(jsonb_build_object('id',id,'tag',task_tag,'status',status,'organization_id',organization_id) ORDER BY task_tag),'[]') FROM public.tasks),
 'month_periods', (SELECT coalesce(jsonb_agg(jsonb_build_object('period',period_start,'state',state,'organization_id',organization_id)),'[]')
   FROM taskpulse_private.month_periods WHERE period_start = ANY({lit('{' + ','.join(months) + '}')}::date[])),
 'monthly_snapshots', (SELECT coalesce(jsonb_agg(to_jsonb(s) ORDER BY id),'[]') FROM public.monthly_performance_snapshots s
   WHERE make_date(snapshot_year,snapshot_month,1) = ANY({lit('{' + ','.join(months) + '}')}::date[])),
 'final_reports', (SELECT coalesce(jsonb_agg(to_jsonb(r) - 'payload' ORDER BY id),'[]') FROM public.report_snapshots r WHERE status='Final'
   AND EXISTS(SELECT 1 FROM unnest({lit('{' + ','.join(months) + '}')}::date[]) m WHERE (r.period_start IS NULL OR r.period_start < (m + interval '1 month')::date)
   AND (r.period_end IS NULL OR r.period_end >= m))),
 'triggers', {TRIGGERS_EXPR},
 'today', (clock_timestamp() AT TIME ZONE 'Asia/Kolkata')::date,
 'receipts', (SELECT coalesce(jsonb_agg(operation_id),'[]') FROM taskpulse_private.command_receipts WHERE command='save_task_with_participants')
)""")


CLOSURE_SQL = """
CREATE TEMP TABLE seed_doomed(rel regclass, id uuid, depth int, root uuid, PRIMARY KEY(rel,id)) ON COMMIT DROP;
CREATE TEMP TABLE seed_leaf(child regclass, col name, parent regclass, deltype "char", PRIMARY KEY(child,col,parent)) ON COMMIT DROP;
DO $closure$
DECLARE fk record; added bigint; n bigint; level int:=0;
BEGIN
 INSERT INTO seed_doomed SELECT 'public.tasks'::regclass, id, 0, id FROM public.tasks WHERE id = ANY(__SEEDS__);
 LOOP
  level:=level+1; added:=0;
  IF level>40 THEN RAISE EXCEPTION 'Dependent-row closure did not converge'; END IF;
  FOR fk IN SELECT c.conrelid::regclass child, c.confrelid::regclass parent, a.attname col, c.confdeltype deltype,
     EXISTS(SELECT 1 FROM pg_attribute i WHERE i.attrelid=c.conrelid AND i.attname='id' AND NOT i.attisdropped AND i.atttypid='uuid'::regtype) has_id
    FROM pg_constraint c JOIN pg_attribute a ON a.attrelid=c.conrelid AND a.attnum=c.conkey[1]
    JOIN pg_attribute p ON p.attrelid=c.confrelid AND p.attnum=c.confkey[1] AND p.attname='id'
    WHERE c.contype='f' AND cardinality(c.conkey)=1 AND c.confdeltype IN ('a','r','c')
     AND c.confrelid IN (SELECT DISTINCT rel FROM seed_doomed)
  LOOP
   IF fk.has_id THEN
    -- Propagate only from the previous level's frontier, so depth is the longest FK path.
    EXECUTE format('INSERT INTO seed_doomed SELECT DISTINCT ON (x.id) %L::regclass, x.id, %s, d.root FROM %s x '
     'JOIN seed_doomed d ON d.rel=%L::regclass AND d.id=x.%I AND d.depth=%s '
     'ON CONFLICT (rel,id) DO UPDATE SET depth=excluded.depth WHERE seed_doomed.depth<excluded.depth', fk.child, level, fk.child, fk.parent, fk.col, level-1);
    GET DIAGNOSTICS n=ROW_COUNT; added:=added+n;
   ELSE
    INSERT INTO seed_leaf VALUES(fk.child, fk.col, fk.parent, fk.deltype) ON CONFLICT DO NOTHING;
   END IF;
  END LOOP;
  EXIT WHEN added=0;
 END LOOP;
END $closure$;
"""

PRIVATE_BY_TASK = ("taskpulse_private.participant_versions", "taskpulse_private.task_credit_balances",
                   "taskpulse_private.earned_credit_events")


def closure_report_expr():
    """jsonb describing the doomed rows (counts, cross-company/cross-task references, full rows)."""
    return """jsonb_build_object(
 'counts', (SELECT coalesce(jsonb_object_agg(rel, n),'{}') FROM (SELECT rel::text rel, count(*) n FROM seed_doomed GROUP BY 1) s),
 'per_task', (SELECT coalesce(jsonb_object_agg(tag, c),'{}') FROM (SELECT t.task_tag tag, jsonb_object_agg(s.rel, s.n) c FROM
   (SELECT root, rel::text rel, count(*) n FROM seed_doomed GROUP BY 1,2) s JOIN public.tasks t ON t.id=s.root GROUP BY 1) q),
 'private', jsonb_build_object(
   'participant_versions', (SELECT count(*) FROM taskpulse_private.participant_versions WHERE task_id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass)),
   'task_credit_balances', (SELECT count(*) FROM taskpulse_private.task_credit_balances WHERE task_id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass)),
   'earned_credit_events', (SELECT count(*) FROM taskpulse_private.earned_credit_events WHERE task_id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass)
     OR log_id IN (SELECT id FROM seed_doomed WHERE rel='public.daily_task_logs'::regclass)),
   'log_revisions', (SELECT count(*) FROM taskpulse_private.log_revisions WHERE log_id IN (SELECT id FROM seed_doomed WHERE rel='public.daily_task_logs'::regclass))),
 'leaf', (SELECT coalesce(jsonb_agg(jsonb_build_object('child',child::text,'col',col,'parent',parent::text,'deltype',deltype)),'[]') FROM seed_leaf),
 'live_logs', (SELECT count(*) FROM public.daily_task_logs WHERE id IN (SELECT id FROM seed_doomed WHERE rel='public.daily_task_logs'::regclass) AND reversed_at IS NULL),
 'foreign_owned', (SELECT coalesce(jsonb_agg(DISTINCT rel::text),'[]') FROM seed_doomed WHERE rel::text = ANY(__FOREIGN__)),
 'extra_tasks', (SELECT coalesce(jsonb_agg(id),'[]') FROM seed_doomed WHERE rel='public.tasks'::regclass AND NOT id = ANY(__SEEDS__)),
 'organizations', (SELECT coalesce(jsonb_agg(DISTINCT organization_id),'[]') FROM public.tasks WHERE id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass)),
 'digest', (SELECT md5(string_agg(rel::text||':'||id::text, ',' ORDER BY rel::text, id)) FROM seed_doomed)
)"""


def rows_expr():
    """Full rows of every doomed and FK-less private row, for the manifest."""
    return """(SELECT jsonb_build_object(
 'doomed', (SELECT coalesce(jsonb_object_agg(rel, rows),'{}') FROM (SELECT d.rel::text rel, jsonb_agg(r.row ORDER BY d.id) rows
    FROM seed_doomed d CROSS JOIN LATERAL (SELECT seed_row(d.rel, d.id) AS row) r GROUP BY 1) s),
 'participant_versions', (SELECT coalesce(jsonb_agg(to_jsonb(x)),'[]') FROM taskpulse_private.participant_versions x WHERE task_id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass)),
 'task_credit_balances', (SELECT coalesce(jsonb_agg(to_jsonb(x)),'[]') FROM taskpulse_private.task_credit_balances x WHERE task_id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass)),
 'earned_credit_events', (SELECT coalesce(jsonb_agg(to_jsonb(x)),'[]') FROM taskpulse_private.earned_credit_events x WHERE task_id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass)
   OR log_id IN (SELECT id FROM seed_doomed WHERE rel='public.daily_task_logs'::regclass)),
 'log_revisions', (SELECT coalesce(jsonb_agg(to_jsonb(x)),'[]') FROM taskpulse_private.log_revisions x WHERE log_id IN (SELECT id FROM seed_doomed WHERE rel='public.daily_task_logs'::regclass))))"""


# A temp function: to_jsonb(row) for (regclass,id) without a persistent object.
ROW_FN = """CREATE FUNCTION pg_temp.seed_row(r regclass, i uuid) RETURNS jsonb LANGUAGE plpgsql AS $f$
DECLARE v jsonb; BEGIN EXECUTE format('SELECT to_jsonb(x) FROM %s x WHERE x.id=$1', r) INTO v USING i; RETURN v; END $f$;
"""


def closure_body(task_ids):
    return CLOSURE_SQL.replace("__SEEDS__", uuid_array(task_ids))


def report_expr(task_ids):
    return closure_report_expr().replace("__SEEDS__", uuid_array(task_ids)).replace(
        "__FOREIGN__", "ARRAY[" + ",".join(lit(t) for t in FOREIGN_OWNED) + "]::text[]")


def snapshot(backend, name, task_ids, with_rows):
    body = closure_body(task_ids) + (ROW_FN if with_rows else "")
    expr = report_expr(task_ids)
    if with_rows:
        expr = f"{expr} || jsonb_build_object('rows', {rows_expr().replace('seed_row(', 'pg_temp.seed_row(')})"
    return backend.query(name, transaction(body, expr))


def delete_sql(task_ids, expected_digest, reason, snapshot_deletes, has_timers):
    """Owner-level transaction: soft-reverse, then delete tasks and every dependent row.
    Only the log/task retain triggers are toggled, plus the timer one when timers exist."""
    toggled = list(GUARDED_TRIGGERS[:2]) + ([GUARDED_TRIGGERS[4]] if has_timers else [])
    # ALTER TABLE refuses tables with pending deferred events (the participant constraint
    # triggers), so flush them first; they stay deferred during the deletes, where a task
    # loses its assignees before the task row itself goes.
    trigger_off = ("SET CONSTRAINTS ALL IMMEDIATE;\n" + "\n".join(f"ALTER TABLE {r} DISABLE TRIGGER {n};" for r, n in toggled)
                   + "\nSET CONSTRAINTS ALL DEFERRED;")
    trigger_on = "SET CONSTRAINTS ALL IMMEDIATE;\n" + "\n".join(f"ALTER TABLE {r} ENABLE TRIGGER {n};" for r, n in toggled)
    snapshot_part = ""
    if snapshot_deletes["monthly"] or snapshot_deletes["reports"]:
        snapshot_part = (
            "ALTER TABLE public.monthly_performance_snapshots DISABLE TRIGGER m07_retain_month;\n"
            "ALTER TABLE public.report_snapshots DISABLE TRIGGER m07_retain_report;\n"
            f"DELETE FROM public.monthly_performance_snapshots WHERE id = ANY({uuid_array(snapshot_deletes['monthly'])});\n"
            f"DELETE FROM public.report_snapshots WHERE id = ANY({uuid_array(snapshot_deletes['reports'])});\n"
            "ALTER TABLE public.monthly_performance_snapshots ENABLE TRIGGER m07_retain_month;\n"
            "ALTER TABLE public.report_snapshots ENABLE TRIGGER m07_retain_report;\n")
    body = closure_body(task_ids) + f"""
SELECT 1 FROM public.tasks WHERE id = ANY({uuid_array(task_ids)}) ORDER BY id FOR UPDATE;
DO $guard$ BEGIN
 IF (SELECT md5(string_agg(rel::text||':'||id::text, ',' ORDER BY rel::text, id)) FROM seed_doomed) IS DISTINCT FROM {lit(expected_digest)} THEN
  RAISE EXCEPTION 'Dependent rows changed since the snapshot; rerun the dry run' USING ERRCODE='40001'; END IF;
END $guard$;
-- Retained performance/presentation history must stop counting these logs before the rows go.
WITH r AS (UPDATE public.daily_task_logs SET reversed_at=clock_timestamp(), execution_version=execution_version+1
  WHERE id IN (SELECT id FROM seed_doomed WHERE rel='public.daily_task_logs'::regclass) AND reversed_at IS NULL RETURNING id, task_id, organization_id)
INSERT INTO taskpulse_private.audit_events(organization_id,actor_id,actor_role,command,table_name,row_id,changed_fields,protected_changes)
 SELECT organization_id, NULL, 'system', 'sheet_seed_soft_reverse', 'daily_task_logs', id, ARRAY['reversed_at'],
  jsonb_build_object('reason', {lit(reason)}, 'task_id', task_id) FROM r;
{trigger_off}
DO $delete$
DECLARE level record; leaf record; n bigint;
BEGIN
 DELETE FROM taskpulse_private.log_revisions WHERE log_id IN (SELECT id FROM seed_doomed WHERE rel='public.daily_task_logs'::regclass);
 DELETE FROM taskpulse_private.earned_credit_events WHERE task_id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass)
  OR log_id IN (SELECT id FROM seed_doomed WHERE rel='public.daily_task_logs'::regclass);
 DELETE FROM taskpulse_private.task_credit_balances WHERE task_id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass);
 DELETE FROM taskpulse_private.participant_versions WHERE task_id IN (SELECT id FROM seed_doomed WHERE rel='public.tasks'::regclass);
 FOR level IN SELECT rel, depth FROM seed_doomed GROUP BY rel, depth ORDER BY depth DESC, rel::text LOOP
  FOR leaf IN SELECT * FROM seed_leaf WHERE parent=level.rel AND deltype IN ('a','r') LOOP
   EXECUTE format('DELETE FROM %s WHERE %I IN (SELECT id FROM seed_doomed WHERE rel=%L::regclass AND depth=%s)', leaf.child, leaf.col, level.rel, level.depth);
  END LOOP;
  EXECUTE format('DELETE FROM %s WHERE id IN (SELECT id FROM seed_doomed WHERE rel=%L::regclass AND depth=%s)', level.rel, level.rel, level.depth);
 END LOOP;
 IF EXISTS(SELECT 1 FROM public.tasks WHERE id = ANY({uuid_array(task_ids)})) THEN RAISE EXCEPTION 'Tasks survived deletion'; END IF;
END $delete$;
{trigger_on}
{snapshot_part}"""
    expr = f"""jsonb_build_object('remaining_tasks', (SELECT count(*) FROM public.tasks WHERE id = ANY({uuid_array(task_ids)})),
 'triggers', {TRIGGERS_EXPR})"""
    return transaction(body, expr)


# --------------------------------------------------------------------------- plan

def plan(args, backend, sheet):
    gates, notes = [], []

    def gate(ok, label, detail=""):
        gates.append((bool(ok), label, detail))
        return ok

    if isinstance(backend, HostedBackend):
        ref = args.project_ref
        gate(len(ref) == 20 and ref.isalpha() and ref.islower(), "project ref is a 20-letter lowercase ref", ref)
        link = args.workdir / "supabase" / ".temp" / "project-ref"
        gate(link.is_file() and link.read_text(encoding="utf-8").strip() == ref, "workdir is linked to that ref", str(link))
        project = next((p for p in backend.projects() if p.get("id") == ref), None)
        gate(project and project.get("status") == "ACTIVE_HEALTHY", "project is ACTIVE_HEALTHY",
             f"{project.get('name') if project else 'not listed'} {project.get('status') if project else ''}")
        if backend.is_production:
            gate(args.confirm_production_ref == PRODUCTION_REF, "production ref confirmed with --confirm-production-ref")
            gate(not args.rehearsal_email_domain, "no --rehearsal-email-domain on production")
            gate(not args.shift_month, "no --shift-month on production")
    else:
        gate(re.fullmatch(r"[a-z][a-z0-9_]{2,62}", args.db) and args.db not in ("postgres", "template0", "template1"),
             "local database name is a disposable name", args.db)

    state = backend.query("state", state_sql(sheet))
    gate(state["ledger"] == args.allow_ledger, f"migration ledger has {args.allow_ledger} entries", f"found {state['ledger']}")
    gate(state["definer_member"], "current user is a member of taskpulse_definer", state["current_user"])
    defaults = [o for o in state["organizations"] if o["is_default"] and o["status"] == "active"]
    org = defaults[0]["id"] if len(defaults) == 1 else None
    gate(org, "exactly one active default company", ", ".join(o["name"] for o in defaults))
    gate(all(m["state"] == "live" for m in state["month_periods"] if m["organization_id"] == org),
         "touched months are absent or live in month_periods",
         ", ".join(f"{m['period']}={m['state']}" for m in state["month_periods"]) or "none")
    today = dt.date.fromisoformat(state["today"])
    future = [f"{t['tag']} {l['log_date']}" for t in sheet["tasks"] for l in t["logs"] if dt.date.fromisoformat(l["log_date"]) > today]
    gate(not future, "no log date is after today (Asia/Kolkata)", ", ".join(future[:5]))

    # People identity matrix.
    users = {u["email"]: u for u in state["users"]}
    employee_ids = {(u["organization_id"], u["employee_id"]): u for u in state["users"]}
    people = []
    for person in sheet["people"]:
        current = users.get(person["email"])
        entry = dict(person)
        if current is None:
            clash = employee_ids.get((org, person["employee_id"]))
            if person["email"] in state["auth_emails"]:
                entry.update(action="conflict", detail="Auth user exists without a profile")
            elif clash:
                entry.update(action="conflict", detail=f"employee id used by {clash['email']}")
            else:
                entry.update(action="create")
        else:
            entry["id"] = current["id"]
            differs = {k: current[c] for k, c in (("name", "full_name"), ("employee_id", "employee_id"), ("role", "role"))
                       if current[c] != person[k]}
            if current["organization_id"] != org or not current["is_active"]:
                entry.update(action="conflict", detail="profile is inactive or in another company")
            elif not differs:
                entry.update(action="exists-skip")
            elif args.update_existing_people or (args.repair and set(differs) == {"role"}):
                clash = employee_ids.get((org, person["employee_id"]))
                if "employee_id" in differs and clash and clash["email"] != person["email"]:
                    entry.update(action="conflict", detail=f"employee id used by {clash['email']}")
                else:
                    entry.update(action="exists-update", before=differs)
            else:
                entry.update(action="conflict", detail=f"differs from the sheet: {differs} (use --update-existing-people)")
        people.append(entry)
    gate(all(p["action"] != "conflict" for p in people), "people identity matrix has no conflicts")
    director = next(p for p in people if p["role"] == "managing_director")
    probe_director = next((u for u in state["users"] if u["role"] == "managing_director" and u["is_active"]
                           and u["organization_id"] == org), None)
    if gate(probe_director, "an existing active managing director to probe impersonation"):
        probe = backend.query("probe", transaction(as_actor(probe_director["id"]) +
            "SELECT set_config('seed.probe', jsonb_build_object('uid',auth.uid(),'role',public.get_user_role(),"
            "'org',public.current_organization())::text, true);",
            "current_setting('seed.probe')::jsonb"))
        gate(probe["uid"] == probe_director["id"] and probe["role"] == "managing_director" and probe["org"] == org,
             "impersonation probe returns auth.uid() and managing_director", f"{probe}")

    # Snapshots that would freeze the imported months for matched people.
    matched_ids = {p.get("id") for p in people if p.get("id")}
    blocking_monthly = [s for s in state["monthly_snapshots"] if s["user_id"] in matched_ids]
    blocking_reports = [r for r in state["final_reports"] if r["employee_id"] is None or r["employee_id"] in matched_ids]
    if blocking_monthly or blocking_reports:
        gate(args.delete_september_snapshots, "imported months have no blocking snapshots (or --delete-september-snapshots)",
             f"{len(blocking_monthly)} monthly, {len(blocking_reports)} Final report snapshot(s)")
    else:
        gate(True, "imported months have no blocking monthly/Final report snapshots")
    gate(all(v == "O" for v in state["triggers"].values()) and len(state["triggers"]) == len(GUARDED_TRIGGERS),
         "guarded triggers are present and enabled", json.dumps(state["triggers"], sort_keys=True))

    # Task identity matrix.
    org_tasks = {t["tag"]: t for t in state["tasks"] if t["organization_id"] == org}
    sheet_tags = {t["tag"] for t in sheet["tasks"]}
    expected = set(args.expect_existing_tags.split(",")) if args.expect_existing_tags else set()
    others = set(org_tasks) - sheet_tags
    deletes = []
    if args.replace_tasks:
        gate(expected and not expected & sheet_tags, "--expect-existing-tags names tasks outside the sheet")
        gate(others <= expected, "every existing non-sheet task is in --expect-existing-tags",
             ", ".join(sorted(others - expected)))
        deletes = sorted(others & expected)
    elif others:
        notes.append(f"{len(others)} existing task(s) outside the sheet are left untouched: {', '.join(sorted(others))}")
    delete_report = None
    if deletes:
        ids = [org_tasks[t]["id"] for t in deletes]
        delete_report = snapshot(backend, "replace-plan", ids, with_rows=False)
        gate(not delete_report["foreign_owned"] and not delete_report["extra_tasks"],
             "tasks to delete have no person-owned or cross-task references",
             f"{delete_report['foreign_owned']} {delete_report['extra_tasks']}")
        gate(delete_report["organizations"] == [org], "tasks to delete belong to the default company")
    tasks = []
    for task in sheet["tasks"]:
        entry = {"tag": task["tag"], "logs": len(task["logs"]), "hours": str(sum(Decimal(l["hours"]) for l in task["logs"]))}
        if task["tag"] in org_tasks:
            entry["action"] = "repair" if args.repair else "exists-skip"
            entry["id"] = org_tasks[task["tag"]]["id"]
        else:
            entry["action"] = "insert"
        tasks.append(entry)
    return {"gates": gates, "notes": notes, "state": state, "org": org, "people": people, "director": director,
            "deletes": deletes, "delete_report": delete_report, "tasks": tasks,
            "snapshots": {"monthly": blocking_monthly, "reports": blocking_reports}}


def print_plan(args, backend, sheet, result, sheet_sha):
    print(f"Target: {backend.label}{' (PRODUCTION)' if backend.is_production else ''}")
    print(f"Sheet: sha256 {sheet_sha}; period {sheet['period']}"
          f"{' shifted to ' + sheet['shifted_period'] if args.shift_month else ''}; {len(sheet['people'])} people, "
          f"{len(sheet['tasks'])} tasks, {sum(len(t['logs']) for t in sheet['tasks'])} logs; progress mode {args.progress_mode}")
    print("\nPre-flight gates:")
    for ok, label, detail in result["gates"]:
        print(f"  {'PASS' if ok else 'FAIL'}  {label}{(' - ' + detail) if detail else ''}")
    print("\nPeople:")
    for p in result["people"]:
        print(f"  {p['action']:<14} {p['email']:<34} {p['employee_id']:<8} {p['role']:<18} {p['name']}"
              f"{(' - ' + p['detail']) if p.get('detail') else ''}{(' - was ' + json.dumps(p['before'])) if p.get('before') else ''}")
    print(f"  director actor: {result['director']['email']}")
    snaps = result["snapshots"]
    if snaps["monthly"] or snaps["reports"]:
        print("\nBlocking snapshots" + (" (will be deleted)" if args.delete_september_snapshots else "") + ":")
        for s in snaps["monthly"]:
            print(f"  monthly_performance_snapshots {s['id']} user {s['user_id']} {s['snapshot_year']}-{s['snapshot_month']:02d}")
        for r in snaps["reports"]:
            print(f"  report_snapshots {r['id']} {r['report_type']} {r['period_start']}..{r['period_end']} employee {r['employee_id']}")
    if result["deletes"]:
        report = result["delete_report"]
        print(f"\nTasks to delete ({len(result['deletes'])}): {', '.join(result['deletes'])}; {report['live_logs']} live log(s) soft-reversed first")
        for tag in result["deletes"]:
            print(f"  {tag}: " + ", ".join(f"{k} {v}" for k, v in sorted(report["per_task"].get(tag, {}).items())))
        print("  FK-less private rows: " + ", ".join(f"{k} {v}" for k, v in report["private"].items()))
        timers = " and public.task_timers.m02_retain_timers" if "public.task_timers" in report["counts"] else ""
        print(f"  one owner transaction; disables public.daily_task_logs.m02_retain_logs, public.tasks.m02_task_guard{timers} "
              "for the deletes and re-enables them (verified tgenabled='O')")
        print("  kept: audit_events, command_receipts, input_history, performance_observations, presentation_history, manual_credit_inputs")
    print("\nTasks:")
    by_tag = {t["tag"]: t for t in sheet["tasks"]}
    for t in result["tasks"]:
        s = by_tag[t["tag"]]
        print(f"  {t['action']:<11} {t['tag']:<12} {s['status']:<11} {s['priority']:<6} scale {s['scale']:>9} "
              f"progress {s['progress']:>6} logs {t['logs']:>3} hours {t['hours']:>7} = page total; "
              f"{s['assignee']} / {', '.join(s['reviewers'])}")
    for note in result["notes"] + sheet["warnings"]:
        print(f"  note: {note}")
    no_log_progress = [t["tag"] for t in sheet["tasks"] if not t["logs"] and Decimal(t["progress"]) > 0]
    if args.progress_mode == "last-log" and no_log_progress:
        print(f"  note: {', '.join(no_log_progress)} carry sheet progress without logs; set with the lifecycle progress command")
    month = sheet["shifted_period"][:7]
    outside = {}
    for t in sheet["tasks"]:
        for log in t["logs"]:
            if log["log_date"][:7] != month:
                outside.setdefault(log["log_date"][:7], []).append(f"{t['tag']} row {log['row']} {log['log_date']}")
    for key in sorted(outside):
        rows = outside[key]
        odd = int(key[:4]) not in (int(month[:4]), int(month[:4]) + 1)
        print(f"  note: {len(rows)} log(s) dated {key} (outside {month}), imported as dated"
              + (f" - CHECK, likely a year typo: {', '.join(rows)}" if odd else ""))
    if args.progress_mode == "last-log":
        print("  note: last-log mode credits each task's sheet progress on its latest-dated log and backdates the "
              "participant share version to the assigned date")


# --------------------------------------------------------------------------- phases

def phase_people(backend, result, manifest):
    created, updated = [], []
    org = result["org"]
    for person in result["people"]:
        if person["action"] == "create":
            uid = backend.create_auth_user(person["email"], {"full_name": person["name"], "employee_id": person["employee_id"]},
                                           {"organization_id": org})
            person["id"] = uid
            created.append({"id": uid, "email": person["email"]})
            manifest["created"]["people"].append({"id": uid, "email": person["email"]})
            manifest["actions"].append(f"created Auth user {person['email']} {uid}")
            save_manifest(manifest)
    statements = []
    for person in result["people"]:
        if person["action"] in ("create", "exists-update"):
            statements.append(f"UPDATE public.users SET role={lit(person['role'])}::public.user_role, full_name={lit(person['name'])},"
                              f" employee_id={lit(person['employee_id'])} WHERE id={lit(person['id'])}::uuid;")
            if person["action"] == "exists-update":
                updated.append({"id": person["id"], "before": person["before"]})
    if statements:
        ids = [p["id"] for p in result["people"]]
        rows = backend.query("people-roles", transaction("\n".join(statements),
            f"(SELECT jsonb_agg(jsonb_build_object('id',id,'email',email,'role',role,'is_active',is_active,'organization_id',organization_id))"
            f" FROM public.users WHERE id = ANY({uuid_array(ids)}))"))
        by_id = {r["id"]: r for r in rows}
        for person in result["people"]:
            row = by_id.get(person["id"])
            if not row or row["role"] != person["role"] or not row["is_active"] or row["organization_id"] != org:
                raise RuntimeError(f"Profile for {person['email']} is not the active {person['role']} in the company: {row}")
        manifest["updated_people"].extend(updated)
        manifest["actions"].append(f"set roles for {len(statements)} profile(s)")
        save_manifest(manifest)
    return {p["email"]: p["id"] for p in result["people"]}


def phase_replace(backend, result, args, manifest):
    snaps = result["snapshots"]
    snapshot_deletes = {"monthly": [s["id"] for s in snaps["monthly"]], "reports": [r["id"] for r in snaps["reports"]]} \
        if args.delete_september_snapshots else {"monthly": [], "reports": []}
    if not result["deletes"] and not snapshot_deletes["monthly"] and not snapshot_deletes["reports"]:
        return
    ids = [t["id"] for t in result["state"]["tasks"] if t["tag"] in result["deletes"] and t["organization_id"] == result["org"]]
    full = snapshot(backend, "replace-snapshot", ids, with_rows=True)
    full_reports = backend.query("snapshot-rows", transaction("", "jsonb_build_object('monthly', (SELECT coalesce(jsonb_agg(to_jsonb(s)),'[]') FROM "
        f"public.monthly_performance_snapshots s WHERE id = ANY({uuid_array(snapshot_deletes['monthly'])})), 'reports', "
        f"(SELECT coalesce(jsonb_agg(to_jsonb(r)),'[]') FROM public.report_snapshots r WHERE id = ANY({uuid_array(snapshot_deletes['reports'])})))"))
    manifest["deleted"] = {"tasks": result["deletes"], "closure": {k: v for k, v in full.items() if k != "rows"},
                           "rows": full.get("rows"), "snapshots": full_reports}
    save_manifest(manifest)
    if ids and full["digest"] != result["delete_report"]["digest"]:
        raise Refusal("Dependent rows changed between the plan and the snapshot; rerun")
    outcome = backend.query("replace", delete_sql(ids, full["digest"], REPLACE_REASON, snapshot_deletes,
                                                       "public.task_timers" in full["counts"])
                            if ids else snapshot_only_sql(snapshot_deletes))
    if outcome.get("remaining_tasks") or any(v != "O" for v in outcome["triggers"].values()):
        raise RuntimeError(f"Replace verification failed: {outcome}")
    manifest["actions"].append(f"deleted {len(ids)} task(s) and dependents; {len(snapshot_deletes['monthly'])} monthly and "
                               f"{len(snapshot_deletes['reports'])} report snapshot(s); triggers {outcome['triggers']}")
    save_manifest(manifest)


def snapshot_only_sql(snapshot_deletes):
    body = ("ALTER TABLE public.monthly_performance_snapshots DISABLE TRIGGER m07_retain_month;\n"
            "ALTER TABLE public.report_snapshots DISABLE TRIGGER m07_retain_report;\n"
            f"DELETE FROM public.monthly_performance_snapshots WHERE id = ANY({uuid_array(snapshot_deletes['monthly'])});\n"
            f"DELETE FROM public.report_snapshots WHERE id = ANY({uuid_array(snapshot_deletes['reports'])});\n"
            "ALTER TABLE public.monthly_performance_snapshots ENABLE TRIGGER m07_retain_month;\n"
            "ALTER TABLE public.report_snapshots ENABLE TRIGGER m07_retain_report;\n")
    return transaction(body, f"jsonb_build_object('remaining_tasks',0,'triggers',{TRIGGERS_EXPR})")


def task_script(task, ids, org, mode, save_op, task_id=None, existing_logs=None):
    """One transaction for one sheet task. task_id is set only in --repair for an existing task."""
    director, assignee = ids[task["director"]], ids[task["assignee"]]
    lines = []
    if task_id is None:
        fields = {"task_tag": task["tag"], "title": task["title"], "category": task["category"], "priority": task["priority"],
                  "due_date": task["due_date"], "task_scale": task["scale"], "comments": task["comments"],
                  "initialize_details": True}
        # Sheet scales are exact binary fractions; float repr round-trips them to the same decimal.
        fields["task_scale"] = float(Decimal(task["scale"]))
        payload = json.dumps(fields, ensure_ascii=False, sort_keys=True)
        lines.append(as_actor(director))
        lines.append(f"SELECT set_config('seed.task', public.save_task_with_participants(NULL,{lit(payload)}::jsonb,"
                     f"{uuid_array([assignee])},{uuid_array([ids[r] for r in task['reviewers']])},ARRAY[1000000],NULL,false,"
                     f"{lit(save_op)}::uuid)::text, true);")
        lines.append(AS_OWNER)
        started = task["assigned_date"] if task["status"] == "In Progress" else None
        lines.append(f"UPDATE public.tasks SET assigned_date={lit(task['assigned_date'])}::date, started_date={lit(started)}::date "
                     "WHERE id=current_setting('seed.task')::uuid;")
        if mode == "last-log":
            lines.append(f"UPDATE taskpulse_private.participant_versions SET effective_at=({lit(task['assigned_date'])}::date::timestamp "
                         "AT TIME ZONE 'Asia/Kolkata') WHERE task_id=current_setting('seed.task')::uuid;")
    else:
        lines.append(f"SELECT set_config('seed.task', {lit(task_id)}, true);")
    logs = task["logs"]
    last = max(range(len(logs)), key=lambda i: (logs[i]["log_date"], i)) if logs else None
    if logs:
        lines.append(as_actor(assignee))
        for index, log in enumerate(logs):
            if existing_logs is not None and log["op"] in existing_logs:
                continue
            delta = task["progress"] if (mode == "last-log" and index == last and task_id is None) else "0"
            payload = {"log_date": log["log_date"], "hours_worked": float(Decimal(log["hours"])),
                       "requested_delta": float(Decimal(delta)), "notes": log["notes"], "raw_links": log["raw_links"]}
            lines.append(f"SELECT public.create_task_log(current_setting('seed.task')::uuid,{jlit(payload)},{lit(log['op'])}::uuid);")
        lines.append(AS_OWNER)
        lines.append("UPDATE public.daily_task_logs SET log_source='Imported' WHERE task_id=current_setting('seed.task')::uuid "
                     "AND log_source<>'Imported';")
    if task_id is None:
        set_progress = Decimal(task["progress"]) > 0 and (mode == "task-only" or not logs)
        steps = []
        if set_progress:
            steps.append((assignee, "progress", {"progress": float(Decimal(task["progress"]))}))
        if task["status"] == "In Progress":
            steps.append((assignee, "transition", {"status": "In Progress"}))
        elif task["status"] == "Discarded":
            steps.append((director, "transition", {"status": "Discarded", "reason": DISCARD_REASON}))
        for actor, command, payload in steps:
            lines.append(AS_OWNER)
            lines.append("SELECT set_config('seed.version', (SELECT lifecycle_version FROM public.tasks "
                         "WHERE id=current_setting('seed.task')::uuid)::text, true);")
            lines.append(as_actor(actor))
            lines.append(f"SELECT public.task_lifecycle_command(current_setting('seed.task')::uuid,{lit(command)},{jlit(payload)},"
                         f"current_setting('seed.version')::bigint,{lit(op_id(save_op, command, payload.get('status', '')))}::uuid);")
    expr = ("(SELECT jsonb_build_object('id',t.id,'tag',t.task_tag,'status',t.status,'progress',t.progress_percentage,"
            "'logs',(SELECT count(*) FROM public.daily_task_logs l WHERE l.task_id=t.id AND l.reversed_at IS NULL),"
            "'hours',(SELECT coalesce(sum(hours_worked),0) FROM public.daily_task_logs l WHERE l.task_id=t.id AND l.reversed_at IS NULL),"
            "'log_ids',(SELECT coalesce(jsonb_agg(id),'[]') FROM public.daily_task_logs l WHERE l.task_id=t.id)) "
            "FROM public.tasks t WHERE t.id=current_setting('seed.task')::uuid)")
    return transaction("\n".join(lines), expr)


def log_keys(logs):
    """Content key per log row (plus its occurrence among identical rows), so inserted or
    reordered sheet rows never re-key the logs already imported."""
    seen = {}
    for log in logs:
        content = f"{log['log_date']}|{log['hours']}|{log['notes'] or ''}|{' '.join(log['raw_links'])}"
        seen[content] = seen.get(content, -1) + 1
        yield log, f"{content}|{seen[content]}"


def log_op(task_key, key, generation):
    return op_id(task_key, "log", key, generation)


def phase_tasks(backend, result, sheet, args, ids, manifest):
    director = result["director"]["email"]
    receipts = set(result["state"]["receipts"])
    by_tag = {t["tag"]: t for t in result["tasks"]}
    for task in sheet["tasks"]:
        entry = by_tag[task["tag"]]
        if entry["action"] == "exists-skip":
            continue
        task = dict(task, director=director)
        if entry["action"] == "insert":
            # Receipts outlive deleted tasks; a fresh generation keeps a re-seed after --undo from replaying one.
            generation = next(g for g in range(100) if op_id(result["org"], task["tag"], "save", g) not in receipts)
            save_op = op_id(result["org"], task["tag"], "save", generation)
            task["logs"] = [dict(l, op=log_op(save_op, key, 0)) for l, key in log_keys(task["logs"])]
            outcome = backend.query(f"task-{task['tag']}", task_script(task, ids, result["org"], args.progress_mode, save_op))
            manifest["created"]["tasks"].append({"id": outcome["id"], "tag": task["tag"], "save_op": save_op,
                                                  "log_ops": [l["op"] for l in task["logs"]], "log_ids": outcome["log_ids"]})
        else:
            candidates = {key: [log_op("?", key, g) for g in range(5)] for _, key in log_keys(task["logs"])}
            existing = backend.query(f"repair-read-{task['tag']}", transaction("", "(SELECT jsonb_build_object('save_op', (SELECT operation_id FROM "
                "taskpulse_private.command_receipts WHERE command='save_task_with_participants' AND result=" + lit(entry["id"]) +
                "::uuid LIMIT 1)))"))
            if not existing["save_op"]:
                raise Refusal(f"{task['tag']}: existing task was not created through save_task_with_participants; repair refused")
            candidates = {key: [log_op(existing["save_op"], key, g) for g in range(5)] for key in candidates}
            all_ops = [o for ops in candidates.values() for o in ops]
            receipts_now = backend.query(f"repair-receipts-{task['tag']}", transaction("", "(SELECT coalesce(jsonb_object_agg(r.operation_id, "
                "EXISTS(SELECT 1 FROM public.daily_task_logs l WHERE l.id=r.result AND l.task_id=" + lit(entry["id"]) + "::uuid AND "
                "l.reversed_at IS NULL)),'{}') FROM taskpulse_private.command_receipts r WHERE r.command='create_log' AND r.operation_id = ANY("
                + uuid_array(all_ops) + "))"))
            logs, missing = [], []
            for log, key in log_keys(task["logs"]):
                ops = candidates[key]
                live = next((o for o in ops if receipts_now.get(o) is True), None)
                fresh = next((o for o in ops if o not in receipts_now), None)
                if live is None and fresh is None:
                    raise Refusal(f"{task['tag']} row {log['row']}: no unused operation id left for repair")
                logs.append(dict(log, op=live or fresh))
                if live is None:
                    missing.append(logs[-1])
            task["logs"] = logs
            if not missing:
                continue
            present = {l["op"] for l in logs} - {l["op"] for l in missing}
            outcome = backend.query(f"repair-{task['tag']}", task_script(task, ids, result["org"], args.progress_mode, None,
                                                                          task_id=entry["id"], existing_logs=present))
            manifest["repaired"].append({"id": entry["id"], "tag": task["tag"], "added_log_ops": [l["op"] for l in missing]})
        if outcome["logs"] < len(task["logs"]) or Decimal(str(outcome["hours"])) != sum(Decimal(l["hours"]) for l in task["logs"]):
            raise RuntimeError(f"{task['tag']}: committed {outcome['logs']} logs / {outcome['hours']} h, sheet has "
                               f"{len(task['logs'])} / {entry['hours']}")
        manifest["actions"].append(f"{entry['action']} {task['tag']} {outcome['id']}: {outcome['logs']} logs, {outcome['hours']} h, "
                                   f"{outcome['status']}, progress {outcome['progress']}")
        save_manifest(manifest)
        print(f"  {entry['action']} {task['tag']}: {outcome['logs']} logs, {outcome['hours']} h, {outcome['status']}, progress {outcome['progress']}")


def phase_verify(backend, result, sheet, args, ids, manifest):
    tags = [t["tag"] for t in sheet["tasks"]]
    director = ids[result["director"]["email"]]
    log_days = sorted(l["log_date"] for t in sheet["tasks"] for l in t["logs"])
    period = dt.date.fromisoformat(sheet["shifted_period"])
    next_month = shift_date(period, 1)
    # The live read is bounded to 62 days, so it covers the sheet month onward; logs dated
    # outside that window (the sheet has one 2006 typo) are expected to be absent from it.
    live_from = period
    live_to = min(max(dt.date.fromisoformat(log_days[-1]) + dt.timedelta(days=1), next_month), period + dt.timedelta(days=62))
    live_expected = sum(1 for d in log_days if live_from.isoformat() <= d < live_to.isoformat())
    body = (as_actor(director) +
            f"SELECT set_config('seed.perf', public.performance_source_context({lit(period)}::date,{lit(next_month)}::date,clock_timestamp())::text, true);\n"
            f"SELECT set_config('seed.live', public.live_presentation_context({lit(live_from)}::date,{lit(live_to)}::date,clock_timestamp(),NULL)::text, true);\n")
    expr = f"""(WITH t AS (SELECT * FROM public.tasks WHERE task_tag = ANY({lit('{' + ','.join(tags) + '}')}::text[]) AND organization_id={lit(result['org'])}::uuid),
 perf AS (SELECT current_setting('seed.perf')::jsonb p), live AS (SELECT current_setting('seed.live')::jsonb l)
 SELECT jsonb_build_object(
 'tasks', (SELECT jsonb_object_agg(t.task_tag, jsonb_build_object('status',t.status,'progress',t.progress_percentage,'scale',t.task_scale,
    'valid',t.valid_for_metrics,'assigned',t.assigned_date,'due',t.due_date,'started',t.started_date,'category',t.category,'priority',t.priority,
    'title',t.title,
    'assignees',(SELECT jsonb_agg(lower(u.email) ORDER BY u.email) FROM public.task_assignees a JOIN public.users u ON u.id=a.user_id WHERE a.task_id=t.id),
    'reviewers',(SELECT jsonb_agg(lower(u.email) ORDER BY u.email) FROM public.task_reviewers r JOIN public.users u ON u.id=r.user_id WHERE r.task_id=t.id),
    'logs',(SELECT coalesce(jsonb_agg(jsonb_build_object('log_date',l.log_date,'hours',l.hours_worked,'notes',l.update_notes,'raw_links',coalesce(to_jsonb(l.raw_links),'[]'))),'[]')
      FROM public.daily_task_logs l WHERE l.task_id=t.id AND l.reversed_at IS NULL AND l.log_source='Imported'),
    'other_logs',(SELECT count(*) FROM public.daily_task_logs l WHERE l.task_id=t.id AND l.reversed_at IS NULL AND l.log_source<>'Imported'),
    'reversed_logs',(SELECT count(*) FROM public.daily_task_logs l WHERE l.task_id=t.id AND l.reversed_at IS NOT NULL))) FROM t),
 'triggers', {TRIGGERS_EXPR},
 'perf_members', (SELECT jsonb_agg(m->>'id') FROM perf, jsonb_array_elements(p->'members') m),
 'perf_task_hours', (SELECT jsonb_object_agg(t.task_tag, x->>'taskHours') FROM perf, jsonb_array_elements(p->'tasks') x JOIN t ON t.id::text=x->>'id'),
 'perf_history', (SELECT p->'historyAvailable' FROM perf),
 'live_logs', (SELECT count(*) FROM live, jsonb_array_elements(l->'logs') g WHERE g->>'task_id' IN (SELECT id::text FROM t)),
 'live_logs_total', (SELECT jsonb_array_length(l->'logs') FROM live),
 'live_available', (SELECT l->'presentationAvailable' FROM live)
))"""
    observed = backend.query("verify", transaction(body, expr))
    failures, digests = [], {}
    by_email = {p["email"]: p for p in sheet["people"]}
    for task in sheet["tasks"]:
        row = observed["tasks"].get(task["tag"]) if observed["tasks"] else None
        if row is None:
            failures.append(f"{task['tag']} missing")
            continue
        expect = {"status": task["status"], "assignees": [task["assignee"]], "reviewers": sorted(task["reviewers"]),
                  "assigned": task["assigned_date"], "due": task["due_date"], "category": task["category"], "priority": task["priority"],
                  "title": task["title"], "started": task["assigned_date"] if task["status"] == "In Progress" else None}
        for key, value in expect.items():
            if row[key] != value:
                failures.append(f"{task['tag']} {key}: {row[key]!r} != {value!r}")
        if Decimal(str(row["progress"])) != Decimal(task["progress"]) or Decimal(str(row["scale"])) != Decimal(task["scale"]):
            failures.append(f"{task['tag']} progress/scale {row['progress']}/{row['scale']} != {task['progress']}/{task['scale']}")
        if row["valid"] != (task["short"] in EXPECTED_VALID):
            failures.append(f"{task['tag']} valid_for_metrics {row['valid']}")
        if row["other_logs"]:
            failures.append(f"{task['tag']} has {row['other_logs']} live non-Imported log(s)")
        if row["reversed_logs"]:
            print(f"  note: {task['tag']} keeps {row['reversed_logs']} reversed log(s) as history")
        db_rows = [dict(r, hours=str(r["hours"])) for r in row["logs"]]
        digests[task["tag"]] = {"sheet": log_digest(task["logs"]), "db": log_digest(db_rows), "logs": len(db_rows),
                                "hours": str(sum(Decimal(str(r["hours"])) for r in db_rows))}
        if digests[task["tag"]]["sheet"] != digests[task["tag"]]["db"]:
            failures.append(f"{task['tag']} log digest differs ({len(db_rows)} logs)")
        perf_hours = (observed["perf_task_hours"] or {}).get(task["tag"])
        if perf_hours is None or Decimal(perf_hours) != sum(Decimal(l["hours"]) for l in task["logs"]):
            failures.append(f"{task['tag']} performance_source_context taskHours {perf_hours}")
    total_logs = sum(d["logs"] for d in digests.values())
    if len(observed["tasks"] or {}) != len(sheet["tasks"]) or total_logs != EXPECTED_TOTAL_LOGS:
        failures.append(f"{len(observed['tasks'] or {})} tasks / {total_logs} live Imported logs")
    if any(v != "O" for v in observed["triggers"].values()) or len(observed["triggers"]) != len(GUARDED_TRIGGERS):
        failures.append(f"triggers {observed['triggers']}")
    members = set(observed["perf_members"] or [])
    people_ids = {ids[e] for e in by_email}
    if not people_ids <= members:
        failures.append(f"performance_source_context members lack {sorted(e for e in by_email if ids[e] not in members)}")
    if observed["live_logs"] != live_expected:
        failures.append(f"live_presentation_context logs for these tasks: {observed['live_logs']} (expected {live_expected})")
    overall = hashlib.sha256("".join(digests[t]["db"] for t in tags if t in digests).encode()).hexdigest()
    valid = sorted(t for t, r in (observed["tasks"] or {}).items() if r["valid"])
    report = {"tasks": len(observed["tasks"] or {}), "live_imported_logs": total_logs,
              "hours": str(sum(Decimal(d["hours"]) for d in digests.values())), "valid_for_metrics": valid,
              "triggers": observed["triggers"], "perf_members": len(members), "sheet_people_in_perf": len(people_ids & members),
              "live_logs": observed["live_logs"], "live_logs_expected": live_expected, "live_window": [str(live_from), str(live_to)],
              "live_logs_total": observed["live_logs_total"],
              "live_presentation_available": observed["live_available"], "digest": overall, "per_task": digests,
              "failures": failures}
    manifest["verification"] = report
    save_manifest(manifest)
    print(f"\nPhase D: {report['tasks']} tasks, {total_logs} live Imported logs, {report['hours']} h; valid {', '.join(valid)}")
    print(f"  performance_source_context: {len(members)} members ({report['sheet_people_in_perf']} sheet people), per-task hours "
          f"{'match' if not any('taskHours' in f for f in failures) else 'DIFFER'}")
    print(f"  live_presentation_context: {observed['live_logs']} logs for these tasks ({observed['live_logs_total']} total; "
          f"{live_expected} of the sheet's logs fall in the window {live_from}..{live_to})")
    print(f"  triggers: {json.dumps(observed['triggers'], sort_keys=True)}")
    print(f"  digest {overall}")
    for tag in tags:
        if tag in digests and digests[tag]["logs"]:
            print(f"    {tag:<12} {digests[tag]['logs']:>3} logs {Decimal(digests[tag]['hours']):>7.2f} h  {digests[tag]['db'][:16]}")
    if failures:
        for failure in failures:
            print(f"  FAIL {failure}")
        raise RuntimeError(f"Phase D verification failed ({len(failures)} issue(s))")
    print("  PASS Phase D verification")


# --------------------------------------------------------------------------- undo

def undo(args, backend, manifest_path):
    source = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    if source["target"] != target_identity(args, backend):
        raise Refusal(f"Manifest target {source['target']} is not this target")
    if source.get("repaired") and not source["created"]["tasks"]:
        raise Refusal("A repair-only manifest cannot be undone on its own; undo the run that created the tasks")
    ids = [t["id"] for t in source["created"]["tasks"]]
    people = source["created"]["people"]
    print(f"Undo {manifest_path}: {len(ids)} task(s), {len(people)} person/people, {len(source['updated_people'])} profile update(s)")
    gates = []
    state = backend.query("undo-state", transaction("", "jsonb_build_object('tasks', (SELECT coalesce(jsonb_agg(id),'[]') FROM public.tasks WHERE id = ANY("
                          f"{uuid_array(ids)})), 'people', (SELECT coalesce(jsonb_agg(id),'[]') FROM public.users WHERE id = ANY("
                          f"{uuid_array([p['id'] for p in people])})), 'ledger', CASE WHEN to_regclass('supabase_migrations.schema_migrations') IS NULL "
                          "THEN NULL ELSE (SELECT count(*) FROM supabase_migrations.schema_migrations) END)"))
    gates.append((state["ledger"] == args.allow_ledger, f"migration ledger has {args.allow_ledger} entries"))
    if backend.is_production:
        gates.append((args.confirm_production_ref == PRODUCTION_REF, "production ref confirmed"))
    report = snapshot(backend, "undo-plan", state["tasks"], with_rows=False) if state["tasks"] else None
    if report:
        gates.append((not report["foreign_owned"] and not report["extra_tasks"], "no person-owned or cross-task references"))
        print("  rows: " + ", ".join(f"{k} {v}" for k, v in sorted(report["counts"].items())))
    for ok, label in gates:
        print(f"  {'PASS' if ok else 'FAIL'}  {label}")
    if not all(ok for ok, _ in gates):
        raise Refusal("Undo gates failed")
    if not args.yes:
        print("Dry run: pass --yes to undo.")
        return
    record = new_manifest(args, backend, source.get("sheet_sha256"))
    record["undo_of"] = str(manifest_path)
    if state["tasks"]:
        full = snapshot(backend, "undo-snapshot", state["tasks"], with_rows=True)
        record["deleted"] = {"closure": {k: v for k, v in full.items() if k != "rows"}, "rows": full["rows"]}
        save_manifest(record)
        outcome = backend.query("undo-delete", delete_sql(state["tasks"], full["digest"], UNDO_REASON, {"monthly": [], "reports": []},
                                                             "public.task_timers" in full["counts"]))
        if outcome["remaining_tasks"] or any(v != "O" for v in outcome["triggers"].values()):
            raise RuntimeError(f"Undo verification failed: {outcome}")
        record["actions"].append(f"deleted {len(state['tasks'])} created task(s)")
    restores = [f"UPDATE public.users SET " + ", ".join(
        f"{ {'name': 'full_name'}.get(k, k)}={lit(v)}{'::public.user_role' if k == 'role' else ''}" for k, v in u["before"].items())
        + f" WHERE id={lit(u['id'])}::uuid;" for u in source["updated_people"]]
    if restores:
        backend.query("undo-people", transaction("\n".join(restores), "to_jsonb(true)"))
        record["actions"].append(f"restored {len(restores)} profile(s)")
    for person in people:
        if person["id"] in state["people"]:
            backend.delete_auth_user(person["id"])
            record["actions"].append(f"deleted Auth user {person['email']}")
    save_manifest(record)
    print(f"Undo complete; manifest {record['path']}")


# --------------------------------------------------------------------------- manifest / main

def target_identity(args, backend):
    if isinstance(backend, HostedBackend):
        return {"kind": "hosted", "ref": args.project_ref}
    return {"kind": "local", "port": args.port, "db": args.db}


def new_manifest(args, backend, sheet_sha):
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    path = args.manifest_dir / f"sheet-seed-{stamp}-{uuid.uuid4().hex[:6]}.json"
    return {"path": str(path), "tool": "tools/seed_sheet_to_project.py", "started_at": stamp, "sheet_sha256": sheet_sha,
            "target": target_identity(args, backend), "options": {k: str(v) for k, v in vars(args).items() if k not in ("psql",)},
            "actions": [], "created": {"people": [], "tasks": []}, "updated_people": [], "repaired": [], "deleted": None}


def save_manifest(manifest):
    path = Path(manifest["path"])
    path.write_text(json.dumps(manifest, indent=1, ensure_ascii=False, default=str), encoding="utf-8")


def download(url, directory):
    if not re.fullmatch(r"https://docs\.google\.com/spreadsheets/d/[\w-]+/export\?format=xlsx", url):
        raise Refusal("--sheet-url must be a Google Sheets xlsx export link")
    directory.mkdir(parents=True, exist_ok=True)
    if any(directory.iterdir()):
        raise Refusal("--download-dir must be a new or empty directory")
    target = directory / "tracker.xlsx"
    with urllib.request.urlopen(url, timeout=60) as response, open(target, "wb") as handle:
        handle.write(response.read())
    return target


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    source = parser.add_mutually_exclusive_group()
    source.add_argument("--xlsx", type=Path)
    source.add_argument("--sheet-url")
    parser.add_argument("--download-dir", type=Path)
    parser.add_argument("--target", choices=("hosted", "local"), required=True)
    parser.add_argument("--project-ref")
    parser.add_argument("--workdir", type=Path)
    parser.add_argument("--confirm-production-ref")
    parser.add_argument("--psql")
    parser.add_argument("--port", type=int, default=15432)
    parser.add_argument("--db")
    parser.add_argument("--dry-run", action="store_true", help="print the plan only (default without --yes)")
    parser.add_argument("--yes", action="store_true", help="apply the plan")
    parser.add_argument("--replace-tasks", action="store_true")
    parser.add_argument("--expect-existing-tags", default="")
    parser.add_argument("--delete-september-snapshots", action="store_true")
    parser.add_argument("--update-existing-people", action="store_true")
    parser.add_argument("--progress-mode", choices=("last-log", "task-only"), default="last-log")
    parser.add_argument("--rehearsal-email-domain")
    parser.add_argument("--repair", action="store_true", help="add only missing logs to tasks this tool created")
    parser.add_argument("--manifest-dir", type=Path, default=Path(tempfile.gettempdir()) / "taskpulse-sheet-seed")
    parser.add_argument("--undo", type=Path, metavar="MANIFEST")
    parser.add_argument("--shift-month", help="rehearsal only: move the sheet month to YYYY-MM")
    parser.add_argument("--allow-ledger", type=int, default=DEFAULT_LEDGER)
    args = parser.parse_args(argv)
    if args.target == "hosted" and not (args.project_ref and args.workdir):
        parser.error("--target hosted needs --project-ref and --workdir")
    if args.target == "local" and not (args.psql and args.db):
        parser.error("--target local needs --psql and --db")
    if not args.undo and not (args.xlsx or args.sheet_url):
        parser.error("give --xlsx or --sheet-url (with --download-dir)")
    if args.sheet_url and not args.download_dir:
        parser.error("--sheet-url needs --download-dir")
    if args.target == "hosted" and args.project_ref == PRODUCTION_REF and (args.rehearsal_email_domain or args.shift_month):
        parser.error("--rehearsal-email-domain and --shift-month are refused on production")
    if args.rehearsal_email_domain and not re.fullmatch(r"[a-z0-9-]+(\.[a-z0-9-]+)+", args.rehearsal_email_domain):
        parser.error("--rehearsal-email-domain must be a domain")
    if args.shift_month and not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", args.shift_month):
        parser.error("--shift-month must be YYYY-MM")
    manifest_dir = args.manifest_dir.resolve()
    if manifest_dir == ROOT or ROOT in manifest_dir.parents:
        parser.error("--manifest-dir must be outside the repository")
    args.manifest_dir = manifest_dir
    return args


def main(argv=None):
    args = parse_args(argv)
    args.manifest_dir.mkdir(parents=True, exist_ok=True)
    sql_dir = args.manifest_dir / f"sql-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:4]}"
    sql_dir.mkdir()
    backend = HostedBackend(sql_dir, args.project_ref, args.workdir) if args.target == "hosted" else \
        LocalBackend(sql_dir, args.psql, args.port, args.db)
    try:
        if args.undo:
            undo(args, backend, args.undo)
            return 0
        path = args.xlsx or download(args.sheet_url, args.download_dir)
        sheet_sha = sha256_file(path)
        months = 0
        if args.shift_month:
            year, month = map(int, args.shift_month.split("-"))
            months = (year * 12 + month) - (2026 * 12 + 9)
        sheet = read_sheet(path, args.rehearsal_email_domain, months)
        result = plan(args, backend, sheet)
        print_plan(args, backend, sheet, result, sheet_sha)
        if not all(ok for ok, _, _ in result["gates"]):
            print("\nREFUSED: one or more gates failed; nothing was written.")
            return 2
        if not args.yes or args.dry_run:
            print("\nDry run: nothing was written. Pass --yes to apply.")
            return 0
        if backend.is_production and args.confirm_production_ref != PRODUCTION_REF:
            raise Refusal("production needs --confirm-production-ref")
        manifest = new_manifest(args, backend, sheet_sha)
        manifest["organization_id"] = result["org"]
        manifest["plan"] = {"people": result["people"], "tasks": result["tasks"], "deletes": result["deletes"]}
        save_manifest(manifest)
        print(f"\nManifest: {manifest['path']}")
        ids = phase_people(backend, result, manifest)
        phase_replace(backend, result, args, manifest)
        phase_tasks(backend, result, sheet, args, ids, manifest)
        phase_verify(backend, result, sheet, args, ids, manifest)
        manifest["finished_at"] = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        save_manifest(manifest)
        return 0
    except Refusal as refusal:
        print(f"\nREFUSED: {refusal}")
        return 2


if __name__ == "__main__":
    sys.exit(main())
