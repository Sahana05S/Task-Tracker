"""Writes tools/fixtures/attendance-september-2026.xlsx: a synthetic billed attendance workbook.

It mirrors the layout of the real "3_1_b_Billed Attendance" month tab (read 22 October 2026):
title row, festival-holiday count, a dates row (one date every four columns from column I),
a STAFF NAME row with weekday labels, an IN / OUT / Worked Hours / Status sub-header row and
one row per staff member. The names and times are made up. Cells mix text (with the trailing
spaces the real export has) and time values, as the real sheet does.

Codes used: P, P L (late), UD, H, AB, F; Sundays and 17 September (a calendar holiday in the
harness) are blank. Run: python -I tools/fixtures/make_attendance_fixture.py
"""
import datetime as dt
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent / "attendance-september-2026.xlsx"
MONTH = dt.date(2026, 9, 1)
DAYS = [MONTH + dt.timedelta(days=i) for i in range(30)]
CALENDAR_BLANK = dt.date(2026, 9, 17)


def times(day, index):
    """Deterministic IN/OUT for a person-day; worked = OUT - IN."""
    start = dt.datetime.combine(day, dt.time(9, 30)) + dt.timedelta(minutes=5 * ((day.day + index) % 4))
    end = dt.datetime.combine(day, dt.time(18, 30)) + dt.timedelta(minutes=10 * ((day.day * 2 + index) % 3))
    return start.time(), end.time(), end - start


def status_for(person, day):
    special = person["special"].get(day.day)
    if special is not None:
        return special
    if day.weekday() == 6 or day == CALENDAR_BLANK:
        return ""
    return person["default"](day)


PEOPLE = [
    # The director's row is all AB, as in the real sheet; the import skips directors.
    {"name": "Meera Vasan", "default": lambda d: "AB", "special": {}},
    {"name": "Arun Kumar S", "default": lambda d: "P L" if d.day % 3 == 0 else "P", "special": {8: "AB", 22: "H"}},
    {"name": "Divya Lakshmi R", "default": lambda d: "P", "special": {15: "UD", 25: "F", 10: "P L"}},
    {"name": "Priya Nair K", "default": lambda d: "P", "special": {}},
]


def build():
    workbook = openpyxl.Workbook()
    key = workbook.active
    key.title = "Attendance key"
    for row in (["", "SETTINGS"], ["", "ATTENDANCE KEY"], ["", "L", "LATE"], ["", "AB", "ABSENT"], ["", "H", "HALF A DAY"],
                ["", "F", "FESTIVAL HOLIDAYS"], ["", "UD", "UNDER HOURS"]):
        key.append(row)
    page = workbook.create_sheet("September 2026")
    page["A1"] = "Synthetic Builders (Test Branch) Billed STAFF ATTENDANCE"
    page["A2"], page["B2"] = "Number of Festival Holidays", 1
    for column, (title, formula) in enumerate([("Total Hours Worked", "Hrs"), ("Cumulative number of Days worked", "P+0.75UD+0.5H"),
                                               ("No. of Days Present", "P"), ("No. of Under-Hour Days", "UD"),
                                               ("No. of Late Days", "L"), ("No. of Full Day Leave", "AB"),
                                               ("No. of Half Day Leave", "H")], start=2):
        page.cell(3, column, title)
        page.cell(4, column, formula)
    page["A4"] = "STAFF NAME"
    for day in DAYS:
        column = 9 + 4 * (day.day - 1)
        page.cell(3, column, day.strftime("%d %b %Y"))
        page.cell(4, column, day.strftime("%a"))
        for k, label in enumerate(["IN", "OUT", "Worked Hours", "Status"]):
            page.cell(5, column + k, label)
    for index, person in enumerate(PEOPLE):
        row = 6 + index
        page.cell(row, 1, person["name"])
        total = dt.timedelta()
        for day in DAYS:
            column = 9 + 4 * (day.day - 1)
            status = status_for(person, day)
            if status in ("P", "P L", "UD", "H"):
                start, end, worked = times(day, index)
                if status == "H":
                    end = (dt.datetime.combine(day, start) + dt.timedelta(hours=4, minutes=15)).time()
                    worked = dt.timedelta(hours=4, minutes=15)
                if status == "UD":
                    end = (dt.datetime.combine(day, start) + dt.timedelta(hours=7, minutes=20)).time()
                    worked = dt.timedelta(hours=7, minutes=20)
                total += worked
                seconds = int(worked.total_seconds())
                text = f"{seconds // 3600}:{seconds % 3600 // 60:02d}:00"
                if day.day % 2:  # text cells, sometimes with a trailing space
                    page.cell(row, column, start.strftime("%H:%M").lstrip("0") + " ")
                    page.cell(row, column + 1, end.strftime("%H:%M") + (" " if day.day % 5 == 0 else ""))
                    page.cell(row, column + 2, text)
                else:  # time values
                    page.cell(row, column, start).number_format = "h:mm"
                    page.cell(row, column + 1, end).number_format = "h:mm"
                    page.cell(row, column + 2, (dt.datetime.min + worked).time()).number_format = "h:mm:ss"
            if status:
                page.cell(row, column + 3, status)
        seconds = int(total.total_seconds())
        page.cell(row, 2, f"{seconds // 3600}:{seconds % 3600 // 60:02d}:00")
    workbook.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
